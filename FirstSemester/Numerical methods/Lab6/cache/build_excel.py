import os
import math
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

def fun(x, y):
    return -0.9 * y

def analytical(x):
    return 0.12 * math.exp(-0.9 * x)

def build_excel():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    y0 = 0.12
    x0 = 0.0
    x1 = 1.0
    n = 20
    h = (x1 - x0) / n
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    
    # ----------------------------------------------------
    # Sheet 1: Метод Ейлера
    # ----------------------------------------------------
    ws_euler = wb.create_sheet(title="Метод Ейлера")
    ws_euler.views.sheetView[0].showGridLines = True
    headers_euler = ["i", "k1", "y1 (Ейлер)", "x", "y (точний)", "Δ1 (похибка)"]
    ws_euler.append(headers_euler)
    
    y_cur = y0
    x_cur = x0
    ws_euler.append([0, "", y0, 0.0, analytical(0.0), 0.0])
    for i in range(1, n + 1):
        k1 = fun(x_cur, y_cur)
        y_cur = y_cur + k1 * h
        x_cur = x_cur + h
        y_exact = analytical(x_cur)
        delta = abs(y_exact - y_cur)
        ws_euler.append([i, k1, y_cur, x_cur, y_exact, delta])
        
    # ----------------------------------------------------
    # Sheet 2: Метод Гюна
    # ----------------------------------------------------
    ws_heun = wb.create_sheet(title="Метод Гюна")
    ws_heun.views.sheetView[0].showGridLines = True
    headers_heun = ["i", "k1", "k2", "y2 (Гюн)", "x", "y (точний)", "Δ2 (похибка)"]
    ws_heun.append(headers_heun)
    
    y_cur = y0
    x_cur = x0
    ws_heun.append([0, "", "", y0, 0.0, analytical(0.0), 0.0])
    for i in range(1, n + 1):
        k1 = fun(x_cur, y_cur)
        k2 = fun(x_cur + h, y_cur + h * k1)
        y_cur = y_cur + (k1 + k2) * h / 2
        x_cur = x_cur + h
        y_exact = analytical(x_cur)
        delta = abs(y_exact - y_cur)
        ws_heun.append([i, k1, k2, y_cur, x_cur, y_exact, delta])

    # ----------------------------------------------------
    # Sheet 3: Метод Рунге - Кутта
    # ----------------------------------------------------
    ws_rk4 = wb.create_sheet(title="Метод Рунге-Кутта")
    ws_rk4.views.sheetView[0].showGridLines = True
    headers_rk4 = ["i", "k1", "k2", "k3", "k4", "y4 (РК4)", "x", "y (точний)", "Δ4 (похибка)"]
    ws_rk4.append(headers_rk4)
    
    y_cur = y0
    x_cur = x0
    ws_rk4.append([0, "", "", "", "", y0, 0.0, analytical(0.0), 0.0])
    for i in range(1, n + 1):
        k1 = fun(x_cur, y_cur)
        k2 = fun(x_cur + h / 2, y_cur + k1 * h / 2)
        k3 = fun(x_cur + h / 2, y_cur + k2 * h / 2)
        k4 = fun(x_cur + h, y_cur + k3 * h)
        y_cur = y_cur + (k1 + 2 * k2 + 2 * k3 + k4) * h / 6
        x_cur = x_cur + h
        y_exact = analytical(x_cur)
        delta = abs(y_exact - y_cur)
        ws_rk4.append([i, k1, k2, k3, k4, y_cur, x_cur, y_exact, delta])

    # ----------------------------------------------------
    # Sheet 4: Зведена таблиця та порівняння
    # ----------------------------------------------------
    ws_comp = wb.create_sheet(title="Порівняння та похибки")
    ws_comp.views.sheetView[0].showGridLines = True
    headers_comp = ["i", "x", "y (точний)", "y1 (Ейлер)", "Δ1", "y2 (Гюн)", "Δ2", "y4 (РК4)", "Δ4"]
    ws_comp.append(headers_comp)
    
    for r in range(2, n + 3):
        # We can extract values from the other sheets
        i_val = ws_euler.cell(row=r, column=1).value
        x_val = ws_euler.cell(row=r, column=4).value
        y_exact_val = ws_euler.cell(row=r, column=5).value
        y1_val = ws_euler.cell(row=r, column=3).value
        delta1_val = ws_euler.cell(row=r, column=6).value
        
        y2_val = ws_heun.cell(row=r, column=4).value
        delta2_val = ws_heun.cell(row=r, column=7).value
        
        y4_val = ws_rk4.cell(row=r, column=6).value
        delta4_val = ws_rk4.cell(row=r, column=9).value
        
        ws_comp.append([i_val, x_val, y_exact_val, y1_val, delta1_val, y2_val, delta2_val, y4_val, delta4_val])

    # Apply styling to all sheets
    for ws in [ws_euler, ws_heun, ws_rk4, ws_comp]:
        # Header style
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
        
        # Data rows style
        for row in range(2, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = data_font
                cell.border = thin_border
                
                # Format numbers
                if col == 1:
                    cell.alignment = center_align
                elif isinstance(cell.value, (int, float)):
                    cell.alignment = right_align
                    val = abs(cell.value)
                    if val != 0 and (val < 1e-4 or val > 1e4):
                        cell.number_format = "0.00000E+00"
                    elif "x" in ws.cell(row=1, column=col).value:
                        cell.number_format = "0.00"
                    else:
                        cell.number_format = "0.000000"

        # Set column widths
        for col in range(1, ws.max_column + 1):
            max_len = max(len(str(ws.cell(row=r, column=col).value or '')) for r in range(1, ws.max_row + 1))
            col_letter = get_column_letter(col)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Add Chart to ws_comp
    chart = LineChart()
    chart.title = "Порівняння чисельних розв'язків з аналітичним y(x)"
    chart.style = 13
    chart.y_axis.title = "y"
    chart.x_axis.title = "x"
    chart.width = 18
    chart.height = 11

    # Data: columns 3 (y exact), 4 (y1), 6 (y2), 8 (y4)
    # Using openpyxl Reference
    data = Reference(ws_comp, min_col=3, min_row=1, max_col=3, max_row=n+2)
    chart.add_data(data, titles_from_data=True)
    for c_idx in [4, 6, 8]:
        d = Reference(ws_comp, min_col=c_idx, min_row=1, max_col=c_idx, max_row=n+2)
        chart.add_data(d, titles_from_data=True)
        
    cats = Reference(ws_comp, min_col=2, min_row=2, max_row=n+2)
    chart.set_categories(cats)
    ws_comp.add_chart(chart, "K2")

    out_path = os.path.join("Lab6", "src", "differential_equation.xlsx")
    wb.save(out_path)
    print(f"Excel file created: {out_path}")

if __name__ == "__main__":
    build_excel()
