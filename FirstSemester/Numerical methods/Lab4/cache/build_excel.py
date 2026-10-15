import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import LineChart, Reference, Series

def solve_systems():
    A = [
        [ 1.0,  1.0,  1.0,  1.0,  1.0],
        [ 0.0,  2.1,  0.0, -0.9, -0.7],
        [-0.5, -0.9,  1.7, -0.4, -0.5],
        [ 0.0, -0.5, -0.5,  2.4, -0.3],
        [-0.2, -0.7, -0.6, -0.6,  1.8]
    ]
    B = [6.0, 0.6, -0.72, 1.32, -0.36]
    X0 = [0.0 for _ in range(5)]
    delta = 1e-10

    dimension = len(X0)
    X = X0[:]
    history = [(0, X[:])]

    for k in range(1, 40):
        for i in range(dimension):
            X[i] = B[i]
            for j in range(dimension):
                if i != j:
                    X[i] -= A[i][j] * X[j]
            X[i] /= A[i][i]

        eps = [abs(X[i] - X0[i]) for i in range(dimension)]
        ksi = [eps[i] / (abs(X[i]) + delta) for i in range(dimension)]
        X0 = X[:]
        history.append((k, X[:]))

        if max(eps) < delta and max(ksi) < delta:
            break

    return history

def create_excel():
    history = solve_systems()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Метод Гауса - Зейделя"

    # Styling elements
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    headers = ["i", "x1", "x2", "x3", "x4", "x5", "x (точний)"]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for k, x_vec in history:
        row_vals = [k] + [round(v, 12) for v in x_vec] + [1.2]
        ws.append(row_vals)

    # Style data rows
    for row in ws.iter_rows(min_row=2, max_row=len(history)+1, min_col=1, max_col=7):
        for idx, cell in enumerate(row):
            cell.font = data_font
            cell.border = thin_border
            if idx == 0:
                cell.alignment = center_align
            else:
                cell.alignment = right_align
                cell.number_format = "0.000000000000" if idx <= 5 else "0.0"

    # Set column widths
    ws.column_dimensions["A"].width = 6
    for col_letter in ["B", "C", "D", "E", "F"]:
        ws.column_dimensions[col_letter].width = 18
    ws.column_dimensions["G"].width = 14

    # Add embedded LineChart for iterations 0 to 10
    # Iteration 0 is row 2, iteration 10 is row 12
    chart = LineChart()
    chart.title = "Зміна поточного наближення розв’язку для методу Гауса — Зейделя (i = 0..10)"
    chart.style = 10
    chart.y_axis.title = "Поточне наближення"
    chart.x_axis.title = "Номер ітерації (i)"
    chart.width = 22
    chart.height = 12

    # Data references: columns B to G (x1..x5, x_точний), rows 1 to 12
    data_ref = Reference(ws, min_col=2, min_row=1, max_col=7, max_row=12)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=12)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)

    # Place chart
    ws.add_chart(chart, "I2")

    output_path = os.path.join("Lab4", "src", "gauss_seidel.xlsx")
    wb.save(output_path)
    print(f"Excel workbook created at: {output_path}")

if __name__ == "__main__":
    create_excel()
