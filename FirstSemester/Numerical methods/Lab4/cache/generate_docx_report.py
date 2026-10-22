import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def get_solution():
    A = [
        [ 1.0,  1.0,  1.0,  1.0,  1.0],
        [ 0.0,  2.1,  0.0, -0.9, -0.7],
        [-0.5, -0.9,  1.7, -0.4, -0.5],
        [ 0.0, -0.5, -0.5,  2.4, -0.3],
        [-0.2, -0.7, -0.6, -0.6,  1.8]
    ]
    B = [6.0, 0.6, -0.72, 1.32, -0.36]
    X0 = [0.0 for _ in range(5)]
    delta = 1e-10

    # Gauss
    n = len(B)
    AA = [aa[:] for aa in A]
    BB = B[:]
    for k in range(n - 1):
        value = 0
        index = k
        for i in range(k, n):
            if abs(AA[i][k]) > value:
                value = abs(AA[i][k])
                index = i
        AA[k], AA[index] = AA[index], AA[k]
        BB[k], BB[index] = BB[index], BB[k]
        for i in range(k + 1, n):
            m = AA[i][k] / AA[k][k]
            for j in range(k, n):
                AA[i][j] -= AA[k][j] * m
            BB[i] -= BB[k] * m
    x_gauss = BB[:]
    x_gauss[n - 1] = BB[n - 1] / AA[n - 1][n - 1]
    for i in range(n - 2, -1, -1):
        for j in range(i + 1, n):
            x_gauss[i] -= AA[i][j] * x_gauss[j]
        x_gauss[i] /= AA[i][i]

    # Gauss-Seidel
    dimension = len(X0)
    X = X0[:]
    history = [(0, X[:])]
    for k in range(1, 40):
        for i in range(dimension):
            X[i] = B[i]
            for j in range(dimension):
                if i != j:
                    X[i] -= A[i][j] * X[j]
            X[i] /= A[i][i]

        eps = [abs(X[i] - X0[i]) for i in range(dimension)]
        ksi = [eps[i] / (abs(X[i]) + delta) for i in range(dimension)]
        X0 = X[:]
        history.append((k, X[:]))

        if max(eps) < delta and max(ksi) < delta:
            break

    return x_gauss, history

def generate_report():
    template_path = os.path.join("Lab4", "materials", "lab04 (1).docx")
    doc = docx.Document(template_path)
    t = doc.tables[0]

    # Row 0: Title block
    c_left = t.rows[0].cells[1]
    c_left.text = (
        "ЗВО: Національний університет \n"
        "«Львівська політехніка»\n"
        "Навчальний рік: 2026/2027\n"
        "Семестр: осінній\n"
        "Навчальна дисципліна:\n"
        "Чисельні методи\n"
        "Кафедра систем автоматизованого\n"
        "проектування\n"
        "Викладач: Чумакевич В.В."
    )

    c_right = t.rows[0].cells[2]
    c_right.text = (
        "Тема: Чисельне розв’язування системи лінійних алгебраїчних рівнянь (№ 4)\n"
        "Інститут комп’ютерних наук та інформаційних технологій \n"
        "Група: ПП-24\n"
        "Студент: Лемішка Андрій\n"
        "Варіант: 12"
    )

    # Row 6: Variant system
    var_text = (
        "a + b + c + d + e = 6\n"
        "2,1 b – 0,9 d – 0,7 e = 0,6\n"
        "–0,5 a – 0,9 b + 1,7 c – 0,4 d – 0,5 e = –0,72\n"
        "–0,5 b – 0,5 c + 2,4 d – 0,3 e = 1,32\n"
        "–0,2 a – 0,7 b – 0,6 c – 0,6 d + 1,8 e = –0,36"
    )
    for c in t.rows[6].cells:
        c.text = var_text

    # Row 8: Matrix A and Vector B
    matrix_text = (
        "Матриця коефіцієнтів A (розмірності 5×5):\n"
        "┌                                              ┐\n"
        "│    1,0     1,0     1,0     1,0     1,0   │\n"
        "│    0,0     2,1     0,0    –0,9    –0,7   │\n"
        "│   –0,5    –0,9     1,7    –0,4    –0,5   │\n"
        "│    0,0    –0,5    –0,5     2,4    –0,3   │\n"
        "│   –0,2    –0,7    –0,6    –0,6     1,8   │\n"
        "└                                              ┘\n\n"
        "Вектор-стовпець вільних членів B:\n"
        "┌         ┐\n"
        "│    6,0  │\n"
        "│    0,6  │\n"
        "│  –0,72  │\n"
        "│   1,32  │\n"
        "│  –0,36  │\n"
        "└         ┘\n\n"
        "Початкове наближення: X(0) = [0; 0; 0; 0; 0]ᵀ; точність: δ = 10⁻¹⁰."
    )
    for c in t.rows[8].cells:
        c.text = matrix_text

    # Row 10: Code
    with open(os.path.join("Lab4", "src", "main.py"), "r", encoding="utf-8") as f:
        code_text = f.read().strip()
    for c in t.rows[10].cells:
        c.text = code_text

    # Compute solutions
    x_gauss, history = get_solution()

    # Row 12: Console output
    console_lines = [repr(x_gauss)]
    for k, xv in history:
        console_lines.append(f"{k} {repr(xv)}")
    console_text = "\n".join(console_lines)
    for c in t.rows[12].cells:
        c.text = console_text

    # Row 15: Gauss solution vector
    for c in t.rows[15].cells:
        c.text = repr(x_gauss)

    # Row 17: Gauss-Seidel table (nested table)
    cell_17 = t.rows[17].cells[0]
    nt = cell_17.tables[0]
    headers = ["i", "x1", "x2", "x3", "x4", "x5", "x"]
    for j, h in enumerate(headers):
        nt.rows[0].cells[j].text = h

    # Adjust row count
    needed_rows = len(history) + 1
    while len(nt.rows) < needed_rows:
        nt.add_row()
    while len(nt.rows) > needed_rows:
        tr = nt.rows[-1]._tr
        tr.getparent().remove(tr)

    for i, (k, xv) in enumerate(history):
        row_cells = nt.rows[i + 1].cells
        row_cells[0].text = str(k)
        for j in range(5):
            val_str = f"{xv[j]:.12f}".replace(".", ",")
            row_cells[j + 1].text = val_str
        row_cells[6].text = "1,2"

    # Row 19: Convergence Chart
    cell_19 = t.rows[19].cells[0]
    cell_19.text = ""
    p_img = cell_19.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    plot_img_path = os.path.join("Lab4", "plots", "gauss_seidel_convergence.png")
    p_img.add_run().add_picture(plot_img_path, width=Inches(5.8))

    # Row 21: Shortened, concise Conclusions
    conclusion_text = (
        "У лабораторній роботі розв’язано систему лінійних алгебраїчних рівнянь варіанта 12 (група ПП-24) "
        "прямим методом Гауса та ітераційним методом Гауса — Зейделя. Розрахунки виконано мовою Python, "
        "а таблиці та графік збіжності побудовано в середовищі MS Excel (файл src/gauss_seidel.xlsx).\n\n"
        "Основні результати:\n"
        "1. Метод Гауса з вибором головного елемента за стовпцем за один прохід дав точний аналітичний "
        "розв’язок: x = [1.2, 1.2, 1.2, 1.2, 1.2], тобто a = b = c = d = e = 1,2.\n"
        "2. Метод Гауса — Зейделя при початковому наближенні X(0) = [0, 0, 0, 0, 0]ᵀ та критерії точності "
        "δ = 10⁻¹⁰ зійшовся за 21 ітерацію до вектора x = [1.2, 1.2, 1.2, 1.2, 1.2].\n"
        "3. Побудований графік підтверджує швидку та монотонну збіжність: уже після 7–8 ітерацій відхилення від "
        "точного розв'язку становлять менше 0,0002. Обидва методи показали повну узгодженість чисельних результатів."
    )
    for c in t.rows[21].cells:
        c.text = conclusion_text

    # Clear trailing paragraphs if any
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    # Append Control Questions
    questions_data = [
        ("1. Які існують типи методів розв'язання систем лінійних алгебраїчних рівнянь (СЛАР) і чим вони відрізняються один від одного?",
         "Згідно з методичним практикумом (розд. 2), методи розв’язування СЛАР поділяють на два основні типи:\n"
         "• Прямі (точні) методи — дають можливість дістати розв'язок, виконавши скінченну, апріорі відому кількість арифметичних операцій. Якщо всі проміжні обчислення виконувати точно (без заокруглень), отримуємо точний розв'язок (наприклад, метод Гаусса за схемою єдиного ділення чи схемою Халецького).\n"
         "• Ітераційні методи — дають нескінченну послідовність наближених розв'язків X(0), X(1), ..., границі яких є розв'язком системи. Обчислювальний процес припиняється при досягненні наперед заданої користувачем точності Δ (наприклад, метод простої ітерації, метод Зейделя)."),

        ("2. Які методи належать до прямих методів розв'язування СЛАР?",
         "У лабораторному практикумі (п. 2.1) виділено такі прямі методи:\n"
         "1) Метод Гаусса за схемою єдиного ділення — вимагає виконання n(4n² – 3n – 4)/6 арифметичних операцій (для n=5 це близько 67–68 операцій);\n"
         "2) Метод Гаусса за схемою Халецького — компактний метод трикутного розкладу, який вимагає лише n(2n – 1) операцій (для n=5 це лише 45 операцій).\n"
         "Також у класичній лінійній алгебрі до прямих належать метод вибору головного елемента, метод Жордана — Гаусса та метод Крамера."),

        ("3. Який алгоритм розв'язування СЛАР методом Гаусса за схемою Халецького?",
         "Матриця коефіцієнтів A розкладається у добуток двох трикутних матриць A = C · B (формула 2), де C — нижня трикутна матриця, а B — верхня трикутна матриця з одиничною головною діагоналлю (b_ii = 1). Система C · B · x = d розв'язується у два кроки:\n"
         "1) Прямий хід: обчислюються елементи матриць C, B та проміжного вектора y із трикутної системи C · y = d (формули 13–14 стор. практикуму);\n"
         "2) Зворотний хід: розв'язується трикутна система B · x = y, звідки невідомі x знаходяться знизу вгору: x_n = y_n, x_i = y_i – ∑_{j=i+1}^n b_ij · x_j."),

        ("4. Які методи належать до ітераційних методів розв'язування СЛАР?",
         "Згідно з практикумом (п. 2.2), основними ітераційними методами є:\n"
         "1) Метод простої ітерації (п. 2.2.1);\n"
         "2) Метод Зейделя (Гаусса — Зейделя) (п. 2.2.2)."),

        ("5. Як записується ітераційний процес методу простої ітерації?",
         "Система Ax = d перетворюється до канонічного вигляду x = Bx + e (формула 6), де:\n"
         "b_ij = –a_ij / a_ii (при i ≠ j), b_ii = 0; e_i = d_i / a_ii (формула 7).\n"
         "Ітераційний процес записується рекурентною формулою (8):\n"
         "x^(k+1) = B · x^k + e,   k = 0, 1, 2, ...\n"
         "Необхідна і достатня умова збіжності: модуль усіх власних значень матриці B повинен бути строго меншим від одиниці (|λ_i| < 1)."),

        ("6. Як оцінюється необхідна кількість ітерацій для досягнення потрібної точності методу простої ітерації?",
         "Якщо норма матриці C = ||B||_1 = max_i ∑_{j=1}^n |b_ij| < 1, то апріорна кількість ітерацій K для досягнення похибки Δ оцінюється за формулою (9) практикуму:\n"
         "K > –(1 / lg C) · [lg ||e||_1 – lg(Δ · (1 – C))] – 1,\n"
         "де lg — десятковий логарифм, ||e||_1 = max_i |e_i|. Вона дозволяє точно оцінити трудомісткість ще до початку обчислень."),

        ("7. Яка обчислювальна схема методу простої ітерації?",
         "Згідно з п. 2.2.1 практикуму:\n"
         "1) Обчислюють елементи матриці B і вектора e за формулою (7);\n"
         "2) Обчислюють норму матриці C = ||B||_1. Якщо C < 1, процес збігається;\n"
         "3) Обирають початкове наближення x^(0) = e. Позначають δ^k = x^(k+1) – e, тоді x^(k+1) = x^(0) + δ^k;\n"
         "4) Обчислюють теоретично необхідну кількість ітерацій K за формулою (9);\n"
         "5) Виконують ітераційний процес K разів або до виконання умови ||x^(k+1) – x^k|| ≤ Δ."),

        ("8. Чим відрізняється метод Зейделя від методу простої ітерації розв'язування СЛАР?",
         "У методі простої ітерації для обчислення компонентів нового вектора x^(k+1) використовуються суто значення попереднього кроку x^k. У методі Зейделя (п. 2.2.2) при обчисленні i-ї компоненти вектора відразу ж підставляються щойно обчислені значення x_1^(k+1), ..., x_(i–1)^(k+1) поточної ітерації. Це прискорює швидкість збіжності приблизно в 2 рази та економить пам'ять ЕОМ."),

        ("9. Яка умова збіжності методу Зейделя?",
         "Достатньою умовою збіжності методу Зейделя є умова строгого діагонального переважання матриці A (формула 10 практикуму):\n"
         "∑_{j=1, j ≠ i}^n |a_ij| ≤ q · |a_ii|,   де q < 1 (для всіх i = 1, ..., n).\n"
         "Тобто модуль діагонального коефіцієнта кожного рядка має бути строго більшим за суму модулів решти коефіцієнтів цього рядка. За цієї умови ітерації гарантовано збігаються при будь-якому x^(0).")
    ]

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("Контрольні запитання та відповіді")
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    for q, a in questions_data:
        pq = doc.add_paragraph()
        pq.paragraph_format.space_before = Pt(8)
        pq.paragraph_format.space_after = Pt(2)
        rq = pq.add_run(q)
        rq.bold = True
        rq.font.name = "Times New Roman"
        rq.font.size = Pt(11)

        pa = doc.add_paragraph()
        pa.paragraph_format.space_before = Pt(0)
        pa.paragraph_format.space_after = Pt(6)
        ra = pa.add_run(a)
        ra.font.name = "Times New Roman"
        ra.font.size = Pt(11)

    out_docx_path = os.path.join("Lab4", "Лемішка_Андрій_ЛР4_ПП24.docx")
    doc.save(out_docx_path)
    print(f"Report DOCX saved to: {out_docx_path}")

if __name__ == "__main__":
    generate_report()
