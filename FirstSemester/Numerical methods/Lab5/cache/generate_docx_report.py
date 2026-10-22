import os
import math
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def fun(x):
    return x**2 + 0.78 * x - 0.108

def dfun(x):
    return 2 * x + 0.78

def run_bisection(a=0.0, b=0.35, delta=1e-10):
    fa = fun(a)
    signa = int(math.copysign(1, fa))
    fb = fun(b)
    signb = int(math.copysign(1, fb))
    history = []
    
    for i in range(60):
        c = (a + b) / 2
        fc = fun(c)
        signc = int(math.copysign(1, fc))
        eps = abs(a - b)
        
        history.append((i, a, signa, b, signb, c, signc))
        
        if eps < delta:
            break
        if signb * signc > 0:
            b = c
            signb = signc
        else:
            a = c
            signa = signc
    return history

def run_newton(x0=0.0, delta=1e-10):
    history = [(0, None, None, x0)]
    x = x0
    for i in range(1, 20):
        f = fun(x)
        df = dfun(x)
        x1 = x - f / df
        eps = abs(x1 - x)
        ksi = 2 * eps / (abs(x1) + delta)
        history.append((i, f, df, x1))
        if eps < delta or ksi < delta:
            break
        x = x1
    return history

def generate_report():
    template_path = os.path.join("Lab5", "materials", "lab05.docx")
    doc = docx.Document(template_path)
    t = doc.tables[0]
    
    # Row 0: Title block
    c_left = t.rows[0].cells[1]
    c_left.text = (
        "ЗВО: Національний університет \n"
        "«Львівська політехніка»\n"
        "Навчальний рік: 2026/2027\n"
        "Семестр: осінній\n"
        "Навчальна дисципліна:\n"
        "Чисельні методи\n"
        "Кафедра систем автоматизованого\n"
        "проектування\n"
        "Викладач: Чумакевич В.В."
    )
    
    c_right = t.rows[0].cells[2]
    c_right.text = (
        "Тема: Чисельне розв’язування нелінійного рівняння (№ 5)\n"
        "Інститут комп’ютерних наук та інформаційних технологій \n"
        "Група: ПП-24\n"
        "Студент: Лемішка Андрій\n"
        "Варіант: 12"
    )
    
    # Row 6: Variant equation
    var_text = "f(x) = x2 + 0,78 x – 0,108"
    for c in t.rows[6].cells:
        c.text = var_text
        
    # Row 9: Analytical roots derivation
    roots_text = (
        "f(x) = x² + 0,78 x – 0,108 = 0\n"
        "Квадратне рівняння ax² + bx + c = 0, де a = 1; b = 0,78; c = –0,108.\n"
        "Дискримінант: D = b² – 4ac = 0,78² – 4 · 1 · (–0,108) = 0,6084 + 0,4320 = 1,0404\n"
        "√D = √1,0404 = 1,02\n\n"
        "Корені рівняння:\n"
        "x₁ = (–0,78 – 1,02) / 2 = –1,80 / 2 = –0,9\n"
        "x₂ = (–0,78 + 1,02) / 2 = 0,24 / 2 = 0,12\n\n"
        "Шуканий додатний корінь: x₂ = 0,12."
    )
    for c in t.rows[9].cells:
        c.text = roots_text
        
    # Row 11: Derivative
    deriv_text = "df(x) / dx = 2 x + 0,78"
    for c in t.rows[11].cells:
        c.text = deriv_text
        
    # Row 13: Program code
    code_path = os.path.join("Lab5", "src", "main.py")
    with open(code_path, "r", encoding="utf-8") as f:
        code_text = f.read().strip()
    for c in t.rows[13].cells:
        c.text = code_text
        
    # Row 15: Console output
    console_path = os.path.join("Lab5", "cache", "console_output.txt")
    with open(console_path, "r", encoding="utf-8") as f:
        console_text = f.read().strip()
    for c in t.rows[15].cells:
        c.text = console_text
        
    # Row 18: Bisection Table (Nested Table 0)
    bis_history = run_bisection()
    cell_18 = t.rows[18].cells[0]
    nt_bis = cell_18.tables[0]
    headers_bis = ["i", "ai", "sign(f(ai))", "bi", "sign(f(bi))", "ci", "sign(f(ci))", "x2"]
    for j, h in enumerate(headers_bis):
        nt_bis.rows[0].cells[j].text = h
        
    needed_rows_bis = len(bis_history) + 1
    while len(nt_bis.rows) < needed_rows_bis:
        nt_bis.add_row()
    while len(nt_bis.rows) > needed_rows_bis:
        tr = nt_bis.rows[-1]._tr
        tr.getparent().remove(tr)
        
    for idx, (i_step, a, sa, b, sb, c_val, sc) in enumerate(bis_history):
        row_cells = nt_bis.rows[idx + 1].cells
        row_cells[0].text = str(i_step)
        row_cells[1].text = f"{a:.12g}".replace(".", ",")
        row_cells[2].text = str(sa)
        row_cells[3].text = f"{b:.12g}".replace(".", ",")
        row_cells[4].text = str(sb)
        row_cells[5].text = f"{c_val:.12g}".replace(".", ",")
        row_cells[6].text = str(sc)
        row_cells[7].text = "0,12"
        
    # Propagate to other cells if merged / separate
    for c in t.rows[18].cells[1:]:
        if len(c.tables) > 0 and c.tables[0] != nt_bis:
            # If there's an unmerged table, update it too
            pass
            
    # Row 20: Newton Table (Nested Table 0)
    newt_history = run_newton()
    cell_20 = t.rows[20].cells[0]
    nt_newt = cell_20.tables[0]
    headers_newt = ["i", "f(xi)", "df(xi)/dx", "xi", "x2"]
    for j, h in enumerate(headers_newt):
        nt_newt.rows[0].cells[j].text = h
        
    needed_rows_newt = len(newt_history) + 1
    while len(nt_newt.rows) < needed_rows_newt:
        nt_newt.add_row()
    while len(nt_newt.rows) > needed_rows_newt:
        tr = nt_newt.rows[-1]._tr
        tr.getparent().remove(tr)
        
    for idx, (i_step, f_val, df_val, x_val) in enumerate(newt_history):
        row_cells = nt_newt.rows[idx + 1].cells
        row_cells[0].text = str(i_step)
        if f_val is None:
            row_cells[1].text = ""
            row_cells[2].text = ""
        else:
            row_cells[1].text = f"{f_val:.15f}".replace(".", ",")
            row_cells[2].text = f"{df_val:.15f}".replace(".", ",")
        row_cells[3].text = f"{x_val:.15f}".replace(".", ",")
        row_cells[4].text = "0,12"
        
    # Row 22: Bisection Plot Image
    cell_22 = t.rows[22].cells[0]
    cell_22.text = ""
    p_img1 = cell_22.paragraphs[0]
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bis_img_path = os.path.join("Lab5", "plots", "bisection_plot.png")
    p_img1.add_run().add_picture(bis_img_path, width=Inches(5.8))
    
    # Row 24: Newton Plot Image
    cell_24 = t.rows[24].cells[0]
    cell_24.text = ""
    p_img2 = cell_24.paragraphs[0]
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    newt_img_path = os.path.join("Lab5", "plots", "newton_plot.png")
    p_img2.add_run().add_picture(newt_img_path, width=Inches(5.8))
    
    # Row 26: Conclusions
    conclusion_text = (
        "У лабораторній роботі досліджено чисельні методи розв’язування нелінійного рівняння варіанта 12 "
        "(група ПП-24): f(x) = x² + 0,78 x – 0,108 = 0. Розрахунки виконано за допомогою складеної програми мовою Python, "
        "а табличні дані та графіки збіжності побудовано в середовищі MS Excel (файл src/nonlinear_equation.xlsx) "
        "та за допомогою бібліотеки Matplotlib.\n\n"
        "Основні результати роботи:\n"
        "1. Аналітичний розв’язок: через обчислення дискримінанта знайдено два дійсні корені x₁ = –0,9 та x₂ = 0,12. "
        "Шуканим додатним коренем квадратного рівняння є x₂ = 0,12.\n"
        "2. Метод половинного ділення (бісекції): на початковому інтервалі ізоляції кореня [0; 0,35] при точності δ = 10⁻¹⁰ "
        "ітераційний процес зійшовся за 33 кроки (ітерації 0..32) до значення x ≈ 0,1200000000. Метод продемонстрував стійку, "
        "але лінійну швидкість звуження інтервалу.\n"
        "3. Метод Ньютона (дотичних): при початковому наближенні x₀ = 0 і тій же точності δ = 10⁻¹⁰ метод показав квадратичну "
        "швидкість збіжності й досяг абсолютно точного значення кореня x = 0,12 уже на 5-й ітерації (|f(x₅)| ~ 10⁻¹⁴).\n"
        "4. Побудовані графіки наочно демонструють стягування меж відрізка до кореня в методі бісекції та надзвичайно швидке "
        "наближення до розв'язку в методі Ньютона. Результати чисельних методів повністю узгодилися з аналітичним розв'язком."
    )
    for c in t.rows[26].cells:
        c.text = conclusion_text
        
    # Clear trailing paragraphs if any
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
        
    # Control Questions & Answers
    questions_data = [
        ("1. Що таке трансцендентне рівняння?",
         "Трансцендентним називається нелінійне рівняння f(x) = 0, що містить трансцендентні функції: тригонометричні (sin, cos, tg), "
         "обернені тригонометричні, показникові (експоненційні) або логарифмічні. На відміну від алгебраїчних рівнянь, корені трансцендентних "
         "рівнянь у загальному випадку не можуть бути виражені через скінченну кількість алгебраїчних операцій (додавання, віднімання, "
         "множення, ділення, добування кореня)."),
         
        ("2. Які чисельні методи використовуються для розв'язання трансцендентних рівнянь?",
         "Для чисельного розв'язування трансцендентних і нелінійних рівнянь застосовують:\n"
         "• Методи локалізації (ізоляції) кореня та бісекції: метод половинного ділення (дихотомії), метод золотого перерізу;\n"
         "• Методи лінійної інтерполяції: метод хорд (пропорційних частин);\n"
         "• Екстраполяційні та дотичні методи: метод Ньютона (метод дотичних), метод січних;\n"
         "• Ітераційні методи: метод простої ітерації (метод релаксації) та його модифікації."),
         
        ("3. Які методи відносяться до методів бісекції?",
         "До методів бісекції (послідовного звуження інтервалу ізоляції кореня [a, b], на кінцях якого функція приймає протилежні знаки "
         "f(a)·f(b) < 0) відносяться: класичний метод половинного ділення (середина відрізка c = (a + b) / 2) та метод золотого січення "
         "(ділення відрізка у пропорції золотого перерізу 0,618)."),
         
        ("4. Як побудований алгоритм методу половинного ділення?",
         "1) Обирається інтервал [a, b], на кінцях якого функція неперервна та f(a)·f(b) < 0.\n"
         "2) Знаходиться середина інтервалу: c = (a + b) / 2, та обчислюється значення функції f(c).\n"
         "3) Якщо |f(c)| < δ або довжина інтервалу |b – a| < δ, обчислення завершуються і c приймається за шуканий корінь.\n"
         "4) Якщо sign(f(b)) · sign(f(c)) > 0 (знаки однакові), то корінь лежить у лівій половині, тому нова права межа b = c.\n"
         "5) Інакше знаки f(a) і f(c) однакові, тоді нова ліва межа a = c.\n"
         "6) Процес ітераційно повторюється з кроку 2."),
         
        ("5. Як оцінюється точність методу половинного ділення?",
         "На кожному кроці довжина інтервалу, що містить точний корінь ξ, зменшується рівно вдвічі. Після N ітерацій довжина становить "
         "(b – a) / 2^N. Гранична абсолютна похибка оцінюється як |ξ – c_N| ≤ (b – a) / 2^(N+1). Кількість ітерацій N, необхідна для досягнення "
         "точності ε, визначається аналітично: N ≥ log₂((b – a) / ε)."),
         
        ("6. Чим відрізняється і які особливості методу золотого січення у порівнянні з методом половинного ділення?",
         "У методі половинного ділення новий поділ інтервалу здійснюється строго посередині (коефіцієнт 0,5), через що на кожному кроці "
         "необхідно обчислювати нове значення функції. У методі золотого січення інтервал ділиться у відношенні золотого перерізу "
         "(τ = (√5 – 1)/2 ≈ 0,618). Завдяки симетрії властивостей золотого перерізу на кожному наступному кроці одна з внутрішніх точок "
         "перевикористовується, що зменшує кількість обчислень трудомістких функцій f(x), забезпечуючи вищу обчислювальну ефективність."),
         
        ("7. Як побудований алгоритм методу хорд?",
         "У методі хорд дуга кривої f(x) на відрізку [a, b] замінюється січною прямою (хордою), що проходить через точки (a, f(a)) та (b, f(b)). "
         "Точка перетину хорди з віссю Ox визначається формулою: x* = a – f(a)·(b – a) / (f(b) – f(a)). "
         "Далі обчислюється f(x*). Залежно від знака f(x*), значенням x* замінюють ту межу, знак функції в якій збігається зі знаком f(x*), "
         "зберігаючи протилежність знаків на кінцях нового інтервалу. Процес повторюється до виконання умови |x*(k+1) – x*(k)| < ε."),
         
        ("8. Як оцінюється похибка методу хорд?",
         "Похибка наближеного розв’язку ξ оцінюється через похідні функції на інтервалі ізоляції [x₁, x₂]: "
         "|ξ – x_n| ≤ (M₁ – m₁) / m₁ · |x_n – x_(n–1)|, де M₁ = max |f'(x)| та m₁ = min |f'(x)| при x ∈ [x₁, x₂]. "
         "Якщо похідна на інтервалі змінюється мало (m₁ ≈ M₁), достатнім критерієм зупинки є |x_n – x_(n–1)| ≤ ε."),
         
        ("9. Як працює метод Ньютона?",
         "Метод Ньютона (метод дотичних) базується на лінеаризації функції f(x) за допомогою розкладу в ряд Тейлора в околі поточної точки x_n: "
         "f(x) ≈ f(x_n) + f'(x_n)·(x – x_n). Прирівнюючи наближення до нуля, отримують ітераційну формулу:\n"
         "x_(n+1) = x_n – f(x_n) / f'(x_n).\n"
         "Геометрично це відповідає проведенню дотичної до графіка y = f(x) у точці (x_n, f(x_n)) до перетину з віссю Ox. "
         "Метод має квадратичну швидкість збіжності в околі простого кореня за умови вдалого вибору початкового наближення: "
         "f(x₀)·f''(x₀) > 0."),
         
        ("10. Як побудований алгоритм методу простої ітерації?",
         "1) Вихідне рівняння f(x) = 0 перетворюється до еквівалентного вигляду x = φ(x) (наприклад, φ(x) = x – λ·f(x)).\n"
         "2) Перевіряється умова збіжності: |φ'(x)| ≤ q < 1 для всіх x з відрізка ізоляції кореня [a, b]. Чим менше q, тим швидша збіжність.\n"
         "3) Обирається довільне початкове наближення x₀ ∈ [a, b].\n"
         "4) Будується рекурентна послідовність наближень: x_(n+1) = φ(x_n),  n = 0, 1, 2, ...\n"
         "5) Ітерації продовжуються до виконання умови зупинки |x_(n+1) – x_n| ≤ (1 – q)/q · ε або |x_(n+1) – x_n| ≤ ε.")
    ]
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("Контрольні запитання та відповіді")
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    for q, a in questions_data:
        pq = doc.add_paragraph()
        pq.paragraph_format.space_before = Pt(8)
        pq.paragraph_format.space_after = Pt(2)
        rq = pq.add_run(q)
        rq.bold = True
        rq.font.name = "Times New Roman"
        rq.font.size = Pt(11)
        
        pa = doc.add_paragraph()
        pa.paragraph_format.space_before = Pt(0)
        pa.paragraph_format.space_after = Pt(6)
        ra = pa.add_run(a)
        ra.font.name = "Times New Roman"
        ra.font.size = Pt(11)
        
    out_docx_path = os.path.join("Lab5", "Лемішка_Андрій_ЛР5_ПП24.docx")
    doc.save(out_docx_path)
    print(f"Report DOCX saved to: {out_docx_path}")

if __name__ == "__main__":
    generate_report()
