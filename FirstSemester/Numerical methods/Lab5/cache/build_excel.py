import os
import math
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import LineChart, Reference

def fun(x):
    return x**2 + 0.78 * x - 0.108

def dfun(x):
    return 2 * x + 0.78

def run_bisection(a=0.0, b=0.35, delta=1e-10):
    fa = fun(a)
    signa = int(math.copysign(1, fa))
    fb = fun(b)
    signb = int(math.copysign(1, fb))
    history = []
    
    for i in range(60):
        c = (a + b) / 2
        fc = fun(c)
        signc = int(math.copysign(1, fc))
        eps = abs(a - b)
        
        history.append((i, a, signa, b, signb, c, signc, 0.12))
        
        if eps < delta:
            break
        if signb * signc > 0:
            b = c
            signb = signc
        else:
            a = c
            signa = signc
    return history

def run_newton(x0=0.0, delta=1e-10):
    history = [(0, None, None, x0, 0.12)]
    x = x0
    for i in range(1, 20):
        f = fun(x)
        df = dfun(x)
        x1 = x - f / df
        eps = abs(x1 - x)
        ksi = 2 * eps / (abs(x1) + delta)
        history.append((i, f, df, x1, 0.12))
        if eps < delta or ksi < delta:
            break
        x = x1
    return history

def create_excel():
    wb = openpyxl.Workbook()
    
    # ----------------- Sheet 1: Bisection -----------------
    ws1 = wb.active
    ws1.title = "Метод половинного ділення"
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    
    headers_bis = ["i", "ai", "sign(f(ai))", "bi", "sign(f(bi))", "ci", "sign(f(ci))", "x2"]
    ws1.append(headers_bis)
    
    for cell in ws1[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        
    bis_history = run_bisection()
    for row in bis_history:
        i, a, sa, b, sb, c, sc, x2 = row
        ws1.append([i, a, sa, b, sb, c, sc, x2])
        
    for row in ws1.iter_rows(min_row=2, max_row=len(bis_history)+1, min_col=1, max_col=8):
        for idx, cell in enumerate(row):
            cell.font = data_font
            cell.border = thin_border
            if idx in [0, 2, 4, 6]:
                cell.alignment = center_align
                cell.number_format = "0"
            else:
                cell.alignment = right_align
                cell.number_format = "0.000000000000" if idx in [1, 3, 5] else "0.00"
                
    ws1.column_dimensions["A"].width = 6
    ws1.column_dimensions["B"].width = 18
    ws1.column_dimensions["C"].width = 14
    ws1.column_dimensions["D"].width = 18
    ws1.column_dimensions["E"].width = 14
    ws1.column_dimensions["F"].width = 18
    ws1.column_dimensions["G"].width = 14
    ws1.column_dimensions["H"].width = 12
    
    # Chart for Bisection (iterations 0 to 10: rows 2 to 12)
    chart1 = LineChart()
    chart1.title = "Зміна меж і поточного наближення кореня (метод половинного ділення, i = 0..10)"
    chart1.style = 10
    chart1.y_axis.title = "Значення"
    chart1.x_axis.title = "Номер ітерації (i)"
    chart1.width = 20
    chart1.height = 12
    
    # Columns B (ai), D (bi), F (ci), H (x2)
    # OpenPyXL LineChart with non-contiguous columns: we add series individually
    for col_idx, col_name in [(2, "ai"), (4, "bi"), (6, "ci"), (8, "x2")]:
        data_ref = Reference(ws1, min_col=col_idx, min_row=1, max_row=12)
        chart1.add_data(data_ref, titles_from_data=True)
    cats_ref1 = Reference(ws1, min_col=1, min_row=2, max_row=12)
    chart1.set_categories(cats_ref1)
    ws1.add_chart(chart1, "J2")
    
    # ----------------- Sheet 2: Newton -----------------
    ws2 = wb.create_sheet(title="Метод Ньютона")
    headers_newt = ["i", "f(xi)", "df(xi)/dx", "xi", "x2"]
    ws2.append(headers_newt)
    
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        
    newt_history = run_newton()
    for row in newt_history:
        i, f, df, x, x2 = row
        ws2.append([i, "" if f is None else f, "" if df is None else df, x, x2])
        
    for row in ws2.iter_rows(min_row=2, max_row=len(newt_history)+1, min_col=1, max_col=5):
        for idx, cell in enumerate(row):
            cell.font = data_font
            cell.border = thin_border
            if idx == 0:
                cell.alignment = center_align
                cell.number_format = "0"
            elif idx == 4:
                cell.alignment = right_align
                cell.number_format = "0.00"
            else:
                cell.alignment = right_align
                if cell.value != "":
                    cell.number_format = "0.000000000000E+00" if idx in [1, 2] else "0.000000000000000"
                    
    ws2.column_dimensions["A"].width = 6
    ws2.column_dimensions["B"].width = 24
    ws2.column_dimensions["C"].width = 20
    ws2.column_dimensions["D"].width = 22
    ws2.column_dimensions["E"].width = 12
    
    # Chart for Newton
    chart2 = LineChart()
    chart2.title = "Зміна поточного наближення кореня (метод Ньютона)"
    chart2.style = 10
    chart2.y_axis.title = "Значення кореня"
    chart2.x_axis.title = "Номер ітерації (i)"
    chart2.width = 18
    chart2.height = 11
    
    data_ref2_x = Reference(ws2, min_col=4, min_row=1, max_row=len(newt_history)+1)
    data_ref2_x2 = Reference(ws2, min_col=5, min_row=1, max_row=len(newt_history)+1)
    chart2.add_data(data_ref2_x, titles_from_data=True)
    chart2.add_data(data_ref2_x2, titles_from_data=True)
    cats_ref2 = Reference(ws2, min_col=1, min_row=2, max_row=len(newt_history)+1)
    chart2.set_categories(cats_ref2)
    ws2.add_chart(chart2, "G2")
    
    out_path = os.path.join("Lab5", "src", "nonlinear_equation.xlsx")
    wb.save(out_path)
    print(f"Excel file created at: {out_path}")

if __name__ == "__main__":
    create_excel()
