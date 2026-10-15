import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

LAB_DIR = r"d:\Studing\SecondYear\FirstSemester\OperatingSystems\Lab5"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy Blue
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x00, 0x4B, 0x87)
    else:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_body_p(doc, text, bold_prefix=None, italic_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        
    if italic_prefix:
        r_it = p.add_run(italic_prefix)
        r_it.italic = True
        r_it.font.name = "Times New Roman"
        r_it.font.size = Pt(12)
        
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_screenshot_figure(doc, img_filename, caption_text, width=Inches(5.8)):
    img_path = os.path.join(LAB_DIR, img_filename)
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_filename} not found!")
        return
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=width)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(10.5)
    run_cap.italic = True
    run_cap.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

print("Helper functions defined successfully")
