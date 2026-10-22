import os
import sys
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

LAB_DIR = r"d:\Studing\SecondYear\FirstSemester\OperatingSystems\Lab5"
REPORTS_DIR = os.path.join(LAB_DIR, "reports")
SCREENSHOTS_DIR = os.path.join(LAB_DIR, "screenshots")

def setup_document_styles(doc):
    sec = doc.sections[0]
    sec.left_margin = Inches(0.984)    # 2.5 cm
    sec.right_margin = Inches(0.591)   # 1.5 cm
    sec.top_margin = Inches(0.591)     # 1.5 cm
    sec.bottom_margin = Inches(0.591)  # 1.5 cm
    sec.page_width = Inches(8.27)      # A4
    sec.page_height = Inches(11.69)
    
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(14)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

def add_title_page(doc):
    # P0: МІНІСТЕРСТВО...
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(1)
    r = p0.add_run("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    # P1: НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ...
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(1)
    r = p1.add_run("НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ «ЛЬВІВСЬКА ПОЛІТЕХНІКА»")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    # P2: Кафедра...
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(6)
    r = p2.add_run("Кафедра «Систем автоматизованого проектування»")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    # P3: Coat of arms (Official Lviv Polytechnic emblem)
    coat_path = os.path.join(SCREENSHOTS_DIR, "lp_coat_of_arms.jpg")
    if os.path.exists(coat_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(4)
        p_logo.paragraph_format.space_after = Pt(10)
        r_logo = p_logo.add_run()
        r_logo.add_picture(coat_path, width=Inches(2.35))

    # P4: Звіт
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(4)
    p4.paragraph_format.space_after = Pt(2)
    r = p4.add_run("Звіт")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(20)

    # P5: до лабораторної роботи №5
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_before = Pt(0)
    p5.paragraph_format.space_after = Pt(1)
    r = p5.add_run("до лабораторної роботи №5")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    # P6: з дисципліни...
    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6.paragraph_format.space_before = Pt(0)
    p6.paragraph_format.space_after = Pt(2)
    r = p6.add_run("з дисципліни «Операційні системи»")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    # P7: на тему:
    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p7.paragraph_format.space_before = Pt(2)
    p7.paragraph_format.space_after = Pt(1)
    r = p7.add_run("на тему:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    # P8: Назва теми
    p8 = doc.add_paragraph()
    p8.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p8.paragraph_format.space_before = Pt(0)
    p8.paragraph_format.space_after = Pt(16)
    r = p8.add_run("«Встановлення та конфігурація служби каталогів Active Directory та служб мережевої інфраструктури (DNS, DHCP, GPO) у середовищі Windows Server»")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    # Submitter block (Right aligned)
    sub_lines = [
        ("Виконав:", False),
        ("студент групи ПП-14", False),
        ("Андрій ЛЕМІШКА", True),
        ("Перевірила:", False),
        ("викладач Наталія НЕСТОР", True)
    ]
    for text, bold in sub_lines:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(1)
        r = p_sub.add_run(text)
        r.bold = bold
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)

    # City & Year (Center aligned at bottom)
    p_city = doc.add_paragraph()
    p_city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_city.paragraph_format.space_before = Pt(20)
    p_city.paragraph_format.space_after = Pt(0)
    r = p_city.add_run("Львів – 2026")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    doc.add_page_break()

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    return p

def add_stage_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = True
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    return p

def add_body_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.first_line_indent = Inches(0.492) # 1.25 cm
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    return p

def add_question(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(f"{num}. {text}")
    r.bold = True
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    return p

def add_screenshot_figure(doc, img_filename, fig_num, caption_text, width=Inches(5.31)):
    img_path = os.path.join(SCREENSHOTS_DIR, img_filename)
    if not os.path.exists(img_path):
        alt_path = os.path.join(LAB_DIR, img_filename)
        if os.path.exists(alt_path):
            img_path = alt_path
        else:
            print(f"Warning: image {img_filename} not found!")
            return
            
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.first_line_indent = Pt(0)
    p_img.paragraph_format.keep_with_next = True
    
    r_img = p_img.add_run()
    r_img.add_picture(img_path, width=width)
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    p_cap.paragraph_format.first_line_indent = Pt(0)
    
    r_lbl = p_cap.add_run(f"Рис. {fig_num} – ")
    r_lbl.bold = True
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(12)
    
    r_txt = p_cap.add_run(caption_text)
    r_txt.italic = True
    r_txt.font.name = "Times New Roman"
    r_txt.font.size = Pt(12)

def build_lab5_report():
    doc = Document()
    setup_document_styles(doc)
    add_title_page(doc)
    
    # 1. МЕТА РОБОТИ
    add_heading_1(doc, "1. МЕТА РОБОТИ")
    add_body_p(doc, "Мета роботи: дослідити особливості інсталяції та адміністрування служби каталогів Active Directory Domain Services (AD DS) у середовищі Windows Server, набути практичних навичок конфігурації базових мережевих служб (DNS, DHCP), управління структурою підрозділів (OU), обліковими записами користувачів і груп, створення об'єктів групової політики безпеки (GPO), організації мережевого файлового доступу (SMB/NTFS) та інтеграції клієнтської робочої станції Windows 10 у домен.")
    
    # 2. ХІД РОБОТИ
    add_heading_1(doc, "2. ХІД РОБОТИ")
    
    # Етап 1
    add_stage_heading(doc, "Етап 1. Налаштування мережевих параметрів сервера та встановлення ролі AD DS")
    add_body_p(doc, "Перед початком розгортання контролера домену виконали конфігурацію мережевого інтерфейсу сервера. Згідно з вимогами завдання серверу призначено статичну IPv4-адресу 192.168.1.5 із маскою підмережі 255.255.255.0 та основним DNS-сервером 127.0.0.1 (петльова адреса для локального DNS).")
    add_screenshot_figure(doc, "screenshot_step1_net_connections.png", 1, "Призначення статичної IP-адреси 192.168.1.5 серверу")
    
    add_body_p(doc, "Відкрили консоль керування «Server Manager» та запустили майстер додавання нових ролей («Add Roles»). На сторінці вибору ролей вибрали службу «Active Directory Domain Services». Майстер автоматично перевірив системні залежності та запропонував додати необхідні компоненти (.NET Framework, оснастки RSAT).")
    add_screenshot_figure(doc, "screenshot_step5_roles_select_adds.png", 2, "Вибір ролі Active Directory Domain Services у диспетчері сервера")
    
    add_body_p(doc, "Підтвердили інсталяцію компонентів та натиснули кнопку «Install». Після успішного копіювання двійкових файлів та службових бібліотек отримали повідомлення про готовність сервера до підвищення ролі за допомогою майстра dcpromo.")
    add_screenshot_figure(doc, "screenshot_step11_install_results.png", 3, "Успішне завершення первинного встановлення ролі AD DS")
    
    # Етап 2
    add_stage_heading(doc, "Етап 2. Підвищення ролі сервера до контролера домену та створення лісу office.com")
    add_body_p(doc, "Через вікно виконання команд (Win + R) або безпосередньо з Server Manager запустили утиліту dcpromo.exe. У стартовому вікні майстра Active Directory Domain Services Installation Wizard ознайомилися з відомостями про сумісність операційних систем.")
    add_screenshot_figure(doc, "screenshot_step12_dcpromo_welcome.png", 4, "Стартове вікно майстра встановлення контролера домену (dcpromo)")
    
    add_body_p(doc, "У вікні конфігурації розгортання обрали створення нового домену в новому лісі («Create a new domain in a new forest») та ввели повне FQDN-ім'я кореневого домену: office.com.")
    add_screenshot_figure(doc, "screenshot_step16_dcpromo_fqdn_office.png", 5, "Введення повного доменного імені кореневого лісу office.com")
    
    add_body_p(doc, "Встановили функціональний рівень лісу та домену на рівні «Windows Server 2008 R2». На кроці додаткових параметрів підтвердили автоматичне встановлення та інтеграцію локального DNS-сервера для обслуговування зони домену.")
    add_screenshot_figure(doc, "screenshot_step18_dcpromo_additional_options.png", 6, "Вибір опції інтегрованого DNS-сервера в майстрі dcpromo")
    
    add_body_p(doc, "Задали шлях розміщення системної бази даних Active Directory (каталог NTDS), файлів журналів транзакцій та загальної папки SYSVOL за замовчуванням (C:\\Windows). Ввели надійний пароль режиму відновлення служб каталогів DSRM (Pa$$word). Після завершення конфігурації сервер виконав автоматичне перезавантаження.")
    add_screenshot_figure(doc, "screenshot_step23_dcpromo_configuring.png", 7, "Процес ініціалізації бази даних NTDS та зон DNS")
    
    # Етап 3
    add_stage_heading(doc, "Етап 3. Побудова ієрархічної структури організаційних підрозділів (OU)")
    add_body_p(doc, "Після перезавантаження виконали вхід під обліковим записом доменного адміністратора (office\\Administrator). Запустили консоль «Active Directory Users and Computers» (dsa.msc). Клацнули правою кнопкою миші на кореневому вузлі office.com -> New -> Organizational Unit та послідовно створили підрозділи: «Керівники», «Фінансисти», «IT-спеціалісти», «Службові» та «Групи».")
    add_screenshot_figure(doc, "screenshot_step30_new_ou_finansysty.png", 8, "Створення організаційного підрозділу «Фінансисти»")
    
    add_body_p(doc, "Перевірили оновлене дерево домену. Усі необхідні структурні підрозділи компанії успішно створено та захищено від випадкового видалення.")
    add_screenshot_figure(doc, "screenshot_step41_tree_all_ous.png", 9, "Структура підрозділів організації в консолі dsa.msc")
    
    # Етап 4
    add_stage_heading(doc, "Етап 4. Створення облікових записів користувачів, груп та налаштування профілів")
    add_body_p(doc, "У відповідних підрозділах створили користувачів відповідно до штатної структури: у підрозділі «Керівники» — користувача «Генеральний директор»; у підрозділі «Фінансисти» — «Фінансовий директор», «Бухгалтер», «Касир»; у підрозділі «IT-спеціалісти» — «Головний IT», «Програміст», «Системний адміністратор». Для кожного облікового запису встановлено пароль «Pa$$word» та прапорець «Password never expires».")
    add_screenshot_figure(doc, "screenshot_step44_ou_finansysty_contents.png", 10, "Перелік облікових записів у підрозділі «Фінансисти»")
    
    add_body_p(doc, "У підрозділі «Групи» створили глобальні групи безпеки «Фінансисти» та «IT-відділ». До групи «Фінансисти» включили користувачів Фінансовий директор, Бухгалтер, Касир. Адміністратора домену включили до групи Domain Admins. Для облікового запису «Програміст» у вкладці «Profile» налаштували підключення мережевої домашньої папки («Connect P: to \\\\server\\Офіс»).")
    add_screenshot_figure(doc, "screenshot_step49_programmer_profile.png", 11, "Налаштування мережевого домашнього диска P: для користувача Програміст")
    
    # Етап 5
    add_stage_heading(doc, "Етап 5. Встановлення, авторизація та запуск служби DHCP-сервера")
    add_body_p(doc, "Через майстер додавання ролей Server Manager встановили роль «DHCP Server». Відкрили оснастку керування DHCP (dhcpmgmt.msc), додали сервер 192.168.1.5, виконали авторизацію сервера в Active Directory та переконалися у запуску системної служби «dhcpserver».")
    add_screenshot_figure(doc, "screenshot_step55_dhcp_service_started.png", 12, "Авторизація та активний стан служби DHCP-сервера в оснастці dhcpmgmt.msc")
    
    # Етап 6
    add_stage_heading(doc, "Етап 6. Створення області DHCP «Офіс» та налаштування пулу адрес")
    add_body_p(doc, "У вузлі IPv4 клацнули правою кнопкою миші -> «New Scope...». Задали ім'я області «Офіс». Вказали початкову IP-адресу 192.168.1.1 та кінцеву 192.168.1.100 із довжиною маски 24 біти (255.255.255.0).")
    add_screenshot_figure(doc, "screenshot_step60_new_scope_ip_range.png", 13, "Задання діапазону розподілу адрес області DHCP")
    
    add_body_p(doc, "На наступному кроці налаштували діапазон виключень (Add Exclusions): початкова адреса 192.168.1.10, кінцева 192.168.1.15. Встановили тривалість оренди адрес (Lease Duration) 8 днів.")
    add_screenshot_figure(doc, "screenshot_step61_new_scope_exclusions.png", 14, "Додавання діапазону адрес виключення 192.168.1.10 – 192.168.1.15")
    
    # Етап 7
    add_stage_heading(doc, "Етап 7. Резервування IP-адреси шлюзу та конфігурація параметрів маршрутизатора і DNS")
    add_body_p(doc, "У майстрі налаштування опцій вказали IP-адресу основного маршрутизатора (Default Gateway) — 192.168.1.4, та адресу DNS-сервера — 192.168.1.5 із доменним іменем office.com. Після активації області перейшли до вузла «Reservations», відкрили діалог «New Reservation» та створили резервування для шлюзу (Reservation name: Gateway, IP: 192.168.1.4, MAC: 1a2b3c4d5e6f).")
    add_screenshot_figure(doc, "screenshot_step68_dhcp_reservations_list.png", 15, "Створене резервування IP-адреси 192.168.1.4 для мережевого шлюзу")
    
    add_body_p(doc, "Перевірили підсумковий стан опцій області DHCP (Scope Options). Сконфігуровано: 003 Router (192.168.1.4), 006 DNS Servers (192.168.1.5) та 015 DNS Domain Name (office.com).")
    add_screenshot_figure(doc, "screenshot_step70_dhcp_scope_options_final.png", 16, "Підсумковий список параметрів області DHCP (опції 003, 006, 015)")
    
    # Етап 8
    add_stage_heading(doc, "Етап 8. Створення та конфігурація групової політики безпеки (GPO) для підрозділу «Фінансисти»")
    add_body_p(doc, "Запустили оснастку «Group Policy Management» (gpmc.msc). Розгорнули вузол домену office.com, клацнули правою кнопкою миші на підрозділі «Фінансисти» та обрали «Create a GPO in this domain, and Link it here...». Вказали ім'я нової політики: «Фінансова політика безпеки».")
    add_screenshot_figure(doc, "screenshot_step71_gpmc_tree.png", 17, "Прив'язка об'єкта групової політики до підрозділу «Фінансисти»")
    
    add_body_p(doc, "Клацнули правою кнопкою миші на політиці -> «Edit...». У редакторі Group Policy Management Editor перейшли до розділу: User Configuration -> Policies -> Administrative Templates -> Control Panel -> Personalization. Відкрили параметр «Screen saver timeout», перевели в стан Enabled, задали значення 600 секунд (10 хвилин) та додали обов'язковий коментар «Фінансова політика безпеки». Також увімкнули параметри «Password protect the screen saver» та «Enable screen saver».")
    add_screenshot_figure(doc, "screenshot_step72_gpo_screensaver_timeout_dialog.png", 18, "Налаштування тайм-ауту блокування екрана та коментаря в редакторі GPO")
    
    # Етап 9
    add_stage_heading(doc, "Етап 9. Організація мережевого спільного доступу до каталогу C:\\Офіс та сховища профілів")
    add_body_p(doc, "У провіднику Windows на диску C: створили каталог «Офіс». У властивостях папки на вкладці «Sharing» відкрили «Advanced Sharing...», увімкнули спільний доступ під іменем «Офіс» та надали права Change і Read для користувачів групи Domain Users. На вкладці «Security» налаштували списки контролю доступу NTFS.")
    add_screenshot_figure(doc, "screenshot_step73_ofis_sharing.png", 19, "Налаштування мережевого доступу SMB для каталогу C:\\Офіс")
    
    add_body_p(doc, "Аналогічно створили папку C:\\Profiles для розміщення профілів користувачів. Надали доступ до мережевого ресурсу \\\\server\\Profiles та встановили права повного доступу для системних служб та адміністраторів.")
    add_screenshot_figure(doc, "screenshot_step75_profiles_sharing.png", 20, "Спільний мережевий доступ до папки розміщення профілів C:\\Profiles")
    
    # Етап 10
    add_stage_heading(doc, "Етап 10. Інтеграція робочої станції Windows 10 у домен та перевірка параметрів DHCP")
    add_body_p(doc, "Запустили клієнтську віртуальну машину Windows 10, підключену до внутрішньої віртуальної мережі office_net. На завантаженому робочому столі відкрили командний рядок та виконали команду ipconfig /all. Клієнтська машина автоматично отримала від DHCP-сервера IP-адресу 192.168.1.1, маску 255.255.255.0, шлюз за замовчуванням 192.168.1.4, DNS-сервер 192.168.1.5 та доменний суфікс office.com.")
    add_screenshot_figure(doc, "screenshot_step77_win10_dhcp_ipconfig.png", 21, "Виведення команди ipconfig /all на робочому столі клієнта Windows 10")
    
    add_body_p(doc, "Відкрили властивості системи (sysdm.cpl), на вкладці «Computer Name» натиснули «Change...», обрали приєднання до домену office.com, ввели облікові дані адміністратора та успішно приєднали робочу станцію до корпоративного домену.")
    
    # Етап 11
    add_stage_heading(doc, "Етап 11. Підвищення привілеїв через механізм Run As та міграція звітів SSRS")
    add_body_p(doc, "Дослідили механізм вторинного входу (Secondary Logon / Run As) для виконання адміністративних завдань без виходу зі звичайного облікового запису користувача. Для запуску консолі керування комп'ютером від імені адміністратора використали команду:")
    add_body_p(doc, "runas /user:office\\Administrator \"mmc compmgmt.msc\"")
    add_body_p(doc, "Також дослідили принципи автоматизованого розгортання та міграції елементів каталогів для служб звітності SQL Server Reporting Services за допомогою утиліти командного рядка rs.exe та скрипта ssrs_migration.rss. Команда виконання міграції має вигляд:")
    add_body_p(doc, "rs.exe -i ssrs_migration.rss -s http://server/reportserver -v SOURCE_URL=\"http://src/reports\" -v TARGET_URL=\"http://dst/reports\"")
    
    # 3. ВІДПОВІДІ НА КОНТРОЛЬНІ ЗАПИТАННЯ (Лаконічні, чіткі відповіді по 2-3 речення)
    add_heading_1(doc, "3. ВІДПОВІДІ НА КОНТРОЛЬНІ ЗАПИТАННЯ")
    
    qa_list = [
        ("Що таке служба каталогів Active Directory та які основні функції вона виконує?",
         "Active Directory Domain Services — це ієрархічна служба каталогів Windows Server для централізованого управління мережевою інфраструктурою. Вона забезпечує єдину базу даних користувачів і ресурсів, аутентифікацію через протокол Kerberos та централізоване застосування групових політик безпеки."),
         
        ("Що таке ліс (Forest), дерево доменів та домен в Active Directory?",
         "Домен є базовою одиницею адміністрування зі спільною базою даних NTDS та політиками безпеки. Дерево доменів об'єднує домени зі спільним неперервним простором імен DNS, а ліс є найвищим рівнем логічної ієрархії зі спільною схемою та глобальним каталогом."),
         
        ("Чим відрізняється робоча група (Workgroup) від домену (Domain)?",
         "У робочій групі кожен комп'ютер є автономним і використовує власну локальну базу облікових записів SAM. У домені база даних є централізованою на контролерах домену, що дозволяє керувати користувачами та правами доступу для всієї мережі з єдиної точки."),
         
        ("Яка роль та призначення служби DNS у доменному середовищі Active Directory?",
         "Служба DNS забезпечує трансляцію доменних імен у IP-адреси та реєстрацію спеціальних SRV-записів. Завдяки SRV-записам клієнтські станції знаходять контролери домену, сервери глобального каталогу та служби автентифікації Kerberos."),
         
        ("Як функціонує служба DHCP та для чого потрібні опції 003, 006 і 015?",
         "DHCP автоматично надає мережеві налаштування клієнтам за протоколом DORA. Опція 003 вказує адресу основного шлюзу маршрутизації, 006 передає перелік IP-адрес DNS-серверів, а 015 задає основний доменний суфікс для комп'ютерів мережі."),
         
        ("Що являють собою групові політики (GPO) та як працює ієрархія їхнього застосування LSDOU?",
         "GPO — це набори параметрів для централізованої конфігурації ОС та прикладного програмного забезпечення. Вони застосовуються в ієрархічному порядку LSDOU: Local -> Site -> Domain -> Organizational Unit, де параметри підрозділів (OU) мають найвищий пріоритет."),
         
        ("Яка різниця між локальними, переміщуваними (Roaming) та обов'язковими (Mandatory) профілями?",
         "Локальний профіль існує лише на одному конкретному ПК. Переміщуваний профіль завантажується з сервера під час входу і синхронізується назад при виході, а обов'язковий профіль (NTUSER.MAN) скидає всі внесені користувачем зміни після завершення сеансу."),
         
        ("Для чого призначена утиліта dcpromo та які компоненти бази даних NTDS вона розгортає?",
         "Утиліта dcpromo підвищує рядовий сервер до ролі контролера домену Active Directory. Вона створює реляційну базу даних NTDS.dit, журнали транзакцій, спільну системну папку SYSVOL та інтегрує необхідні зони DNS."),
         
        ("Як працює механізм Run As для тимчасового підвищення адміністративних привілеїв?",
         "Механізм Run As реалізує вторинний вхід (Secondary Logon), дозволяючи користувачеві запускати окремі програми або консолі з маркером безпеки адміністратора. Це забезпечує принцип найменших привілеїв без постійної роботи під адмін-акаунтом."),
         
        ("Яким чином здійснюється автоматизоване розгортання та міграція звітів у службах SSRS?",
         "Автоматизована міграція звітів SQL Server Reporting Services виконується утилітою командного рядка rs.exe та скриптами .rss. Вона через SOAP API підключається до репозиторію звітів, експортує визначення RDL і джерела даних та публікує їх на цільовому сервері.")
    ]
    
    for i, (q, a) in enumerate(qa_list, 1):
        add_question(doc, i, q)
        add_body_p(doc, a)
        
    # 4. ВИСНОВКИ
    add_heading_1(doc, "4. ВИСНОВКИ")
    add_body_p(doc, "Виконуючи лабораторну роботу №5, я успішно дослідив принципи проєктування, розгортання та централізованого управління корпоративною мережевою інфраструктурою на базі операційної системи Microsoft Windows Server та клієнтської станції Windows 10.")
    add_body_p(doc, "У процесі практичної роботи набув навичок інсталяції та налаштування служби каталогів Active Directory Domain Services (AD DS) із створенням нового кореневого домену office.com за допомогою майстра dcpromo. Побудував ієрархічну структуру організаційних підрозділів (OU) за функціональним призначенням, створив облікові записи співробітників компанії та глобальні групи безпеки для розмежування прав.")
    add_body_p(doc, "Налаштував та авторизував у домені службу DHCP-сервера: створив пул розподілу динамічних адрес «Офіс», налаштував виключення адрес, резервування IP-адреси шлюзу за MAC-адресою та конфігурацію опцій маршрутизатора і DNS. Дослідив механізми застосування об'єктів групової політики (GPO), налаштувавши політику безпеки блокування екрана із захистом паролем для фінансового відділу.")
    add_body_p(doc, "Організував спільні мережеві ресурси (SMB/NTFS), підключив мережевий диск для користувача та успішно інтегрував клієнтську операційну систему Windows 10 у домен office.com, підтвердивши отримання всіх параметрів мережі від сервера.")
    
    out_docx = os.path.join(REPORTS_DIR, "Лемішка_ПП_14_Лаб_5.docx")
    doc.save(out_docx)
    print(f"Report successfully saved to {out_docx}")

if __name__ == "__main__":
    build_lab5_report()
