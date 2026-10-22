import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc1_path = r"D:\Studing\SecondYear\FirstSemester\OperatingSystems\Lab1\Лемішка_ПП_14_Лаб_1.docx"
doc = docx.Document(doc1_path)

print("Images in Lab1:")
for i, p in enumerate(doc.paragraphs):
    if "graphic" in p._p.xml:
        cap = doc.paragraphs[i+1].text if i+1 < len(doc.paragraphs) else "No caption"
        print(f"P{i} [align={p.alignment}] has image -> Caption (P{i+1}): {cap}")

print("\nAll Headings/Sections in Lab1:")
for i, p in enumerate(doc.paragraphs):
    if p.runs and any(r.bold for r in p.runs) and len(p.text.strip()) > 0 and len(p.text.strip()) < 100:
        if not p.text.startswith("Рис."):
            print(f"H P{i:02d} [{p.alignment}]: {p.text}")
