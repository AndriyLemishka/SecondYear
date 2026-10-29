import os
import sys
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import win32com.client

LAB_DIR = r"d:\Studing\SecondYear\FirstSemester\OperatingSystems\Lab6"
REPORTS_DIR = os.path.join(LAB_DIR, "reports")
SCREENSHOTS_DIR = os.path.join(LAB_DIR, "screenshots")
os.makedirs(REPORTS_DIR, exist_ok=True)

COLOR_PRIMARY = RGBColor(0x1B, 0x36, 0x5D)    # Navy Blue
COLOR_SECONDARY = RGBColor(0x00, 0x4B, 0x87)  # Dark Blue
COLOR_TEXT = RGBColor(0x22, 0x22, 0x22)       # Charcoal Body Text
COLOR_MUTED = RGBColor(0x55, 0x55, 0x55)      # Caption Gray

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def setup_document_styles(doc):
    sec = doc.sections[0]
    sec.left_margin = Inches(0.984)    # 2.5 cm
    sec.right_margin = Inches(0.591)   # 1.5 cm
    sec.top_margin = Inches(0.591)     # 1.5 cm
    sec.bottom_margin = Inches(0.591)  # 1.5 cm
    sec.page_width = Inches(8.27)      # A4
    sec.page_height = Inches(11.69)
    
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = COLOR_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

def add_header(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = COLOR_PRIMARY
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_SECONDARY
    else:
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_TEXT
    return p

def add_p(doc, text="", bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.color.rgb = COLOR_TEXT
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(12)
        r_b.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_TEXT
    return p

def add_image_figure(doc, img_filename, caption_text):
    img_path = os.path.join(SCREENSHOTS_DIR, img_filename)
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found!")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.keep_with_next = True
    r_img = p_img.add_run()
    r_img.add_picture(img_path, width=Inches(5.6))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)
    r_cap.italic = True
    r_cap.font.color.rgb = COLOR_MUTED

def add_title_page(doc, doc_type="ЗВІТ", lab_num="6", theme="«Віддалений доступ в Linux»"):
    # Header org
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(1)
    r = p0.add_run("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(1)
    r = p1.add_run("НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ «ЛЬВІВСЬКА ПОЛІТЕХНІКА»")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(6)
    r = p2.add_run("Кафедра «Систем автоматизованого проектування»")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    coat_path = os.path.join(SCREENSHOTS_DIR, "lp_coat_of_arms.jpg")
    if os.path.exists(coat_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(4)
        p_logo.paragraph_format.space_after = Pt(8)
        r_logo = p_logo.add_run()
        r_logo.add_picture(coat_path, width=Inches(2.2))

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(4)
    p4.paragraph_format.space_after = Pt(2)
    r = p4.add_run(doc_type)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(18)
    r.font.color.rgb = COLOR_PRIMARY

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_before = Pt(0)
    p5.paragraph_format.space_after = Pt(1)
    r = p5.add_run(f"до лабораторної роботи №{lab_num}")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6.paragraph_format.space_before = Pt(0)
    p6.paragraph_format.space_after = Pt(2)
    r = p6.add_run("з дисципліни «Операційні системи»")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p7.paragraph_format.space_before = Pt(2)
    p7.paragraph_format.space_after = Pt(1)
    r = p7.add_run("на тему:")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    p8 = doc.add_paragraph()
    p8.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p8.paragraph_format.space_before = Pt(0)
    p8.paragraph_format.space_after = Pt(18)
    r = p8.add_run(theme)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_PRIMARY

    sub_lines = [
        ("Виконав:", False),
        ("студент групи ПП-14", False),
        ("Лемішка Андрій", True),
        ("", False),
        ("Перевірив:", False),
        ("викладач кафедри САПР", False),
        ("Белей О. І.", True)
    ]
    for txt, is_b in sub_lines:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(1)
        p_sub.paragraph_format.right_indent = Inches(0.4)
        if txt:
            r_sub = p_sub.add_run(txt)
            r_sub.font.name = "Times New Roman"
            r_sub.font.size = Pt(12)
            r_sub.bold = is_b

    p_city = doc.add_paragraph()
    p_city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_city.paragraph_format.space_before = Pt(28)
    p_city.paragraph_format.space_after = Pt(0)
    r_city = p_city.add_run("Львів – 2026")
    r_city.bold = True
    r_city.font.name = "Times New Roman"
    r_city.font.size = Pt(12)
    
    doc.add_page_break()

# ==========================================================
# BUILD DOCUMENT 1: MAIN REPORT (Лемішка_ПП_14_Лаб_6.docx)
# ==========================================================
def build_report_doc():
    print("Building Main Academic Report: Лемішка_ПП_14_Лаб_6.docx...")
    doc = Document()
    setup_document_styles(doc)
    add_title_page(doc, doc_type="ЗВІТ", lab_num="6", theme="«Віддалений доступ в Linux»")

    # Header section
    add_header(doc, "Лабораторна робота №6", 1)
    add_header(doc, "Тема: Віддалений доступ в Linux", 2)
    
    add_p(doc, "Ознайомитися на практиці із засобами віддаленого управління в операційній системі Linux Ubuntu. Набути досвіду і навички управління віддаленим доступом за допомогою протоколів SSH та SFTP, налаштування служби sshd, взаємодії з Windows-клієнтом через пакет PuTTY, використання асиметричних ключів шифрування, організації тунелювання TCP-портів та графічного віддаленого доступу.", bold_prefix="Мета роботи: ")
    
    add_p(doc, bold_prefix="Обладнання та програмне забезпечення:")
    add_bullet(doc, "Персональний комп'ютер під управлінням ОС Microsoft Windows 11 Pro;", bold_prefix="1. Хост-система: ")
    add_bullet(doc, "Середовище віртуалізації Oracle VM VirtualBox 7.x із конфігурацією віртуальної машини Ubuntu_Lab6;", bold_prefix="2. Віртуалізація: ")
    add_bullet(doc, "Дистрибутив Linux Ubuntu Server / Desktop 64-bit із мережевими адаптерами NAT та Host-Only Network;", bold_prefix="3. Гостьова ОС: ")
    add_bullet(doc, "Пакет OpenSSH (демон sshd версії 10.2p1, клієнт ssh, scp, sftp);", bold_prefix="4. Серверні служби: ")
    add_bullet(doc, "Комплект клієнтського ПЗ PuTTY Suite версії 0.85 (putty.exe, puttygen.exe, pageant.exe, pscp.exe, psftp.exe, plink.exe);", bold_prefix="5. Клієнтське ПЗ Windows: ")
    add_bullet(doc, "Служби віддаленого робочого столу Remote Desktop (mstsc / xrdp) та X2Go Client.", bold_prefix="6. Графічний доступ: ")

    add_header(doc, "1. Теоретичні відомості", 1)
    
    add_header(doc, "1.1. Порівняльний аналіз протоколів віддаленого доступу: Telnet vs SSH", 2)
    add_p(doc, "Операційні системи сімейства UNIX спочатку проектувалися з орієнтацією на мережеву взаємодію у режимі клієнт-сервер. Історично першим широковідомим стандартом віддаленого термінального управління був протокол Telnet (порт 23 TCP), розроблений для мережі ARPANET. Однак Telnet володіє критичними архітектурними вразливостями: відсутністю криптографічного захисту (усі дані, включно з обліковими записами та паролями, передаються у відкритому текстовому вигляді Plaintext) та відсутністю взаємної автентифікації хостів, що робить його вразливим до атак типу Man-in-the-Middle (MITM) та перехоплення трафіку сниферами (Wireshark, tcpdump).")
    add_p(doc, "Протокол SSH (Secure Shell, порт 22 TCP) розроблено як повноцінну, криптографічно захищену заміну застарілим сервісам Telnet, rlogin та rsh. SSH забезпечує конфіденційність, цілісність даних та строгу автентифікацію як користувача, так і віддаленого сервера. Сучасний стандарт SSH-2 розділений на три незалежні протоколи:")
    add_bullet(doc, "забезпечує шифрування, перевірку цілісності за допомогою HMAC та узгодження ключів сесії за алгоритмом Діффі-Хеллмана або на еліптичних кривих (ECDH, Curve25519);", bold_prefix="• Транспортний рівень (Transport Layer): ")
    add_bullet(doc, "відділений від транспорту та підтримує автентифікацію за паролем, відкритим ключем (RSA, ECDSA, Ed25519), Kerberos або сертифікатами;", bold_prefix="• Протокол автентифікації (User Authentication Protocol): ")
    add_bullet(doc, "мультиплексує один шифрований тунель на множину логічних каналів (інтерактивні термінали, перенаправлення TCP-портів Port Forwarding, X11 Forwarding).", bold_prefix="• Протокол з'єднання (Connection Protocol): ")

    add_header(doc, "1.2. Протоколи передачі файлів: SFTP vs FTP", 2)
    add_p(doc, "Класичний протокол FTP передає команди та вміст файлів відкритим текстом, створюючи ризик компрометації облікових даних. На відміну від FTP, SFTP (SSH File Transfer Protocol) не є окремим сервером, а працює як підсистема (Subsystem sftp) всередині захищеного каналу SSH-2. SFTP шифрує всі транзакції, не вимагає відкриття додаткових динамічних портів у брандмауері (достатньо стандартного порту 22) та гарантує повну безпеку маніпуляцій із файловою системою.")

    add_header(doc, "2. Хід практичного виконання роботи", 1)

    # Step 1: VBox creation
    add_header(doc, "2.1. Створення та конфігурація віртуальної машини Ubuntu у VirtualBox", 2)
    add_p(doc, "Згідно з методичними рекомендаціями викладача (Лаба6_Рекоменд.pdf), на початковому етапі розгортається віртуальна машина Ubuntu у середовищі VirtualBox. Задано назву машини Ubuntu_Lab6, тип операційної системи Linux, версію Ubuntu (64-bit), виділено 2048 МБ оперативної пам'яті (RAM), 2 процесорні ядра та динамічний віртуальний накопичувач формату VDI обсягом 20 ГБ.")
    add_image_figure(doc, "fig01_vbox_create_wizard.png", "Рис. 1. Майстер створення віртуальної машини Ubuntu_Lab6 у VirtualBox")
    add_image_figure(doc, "fig02_vbox_ram_and_disk.png", "Рис. 2. Налаштування обсягу оперативної пам'яті (2048 MB) та накопичувача VDI (20 GB)")

    # Step 2: Networking
    add_header(doc, "2.2. Налаштування мережевих інтерфейсів (NAT та Host-Only Network)", 2)
    add_p(doc, "Для забезпечення ізольованого, стабільного та захищеного підключення хоста Windows до віртуального сервера Linux сконфігуровано два мережеві адаптери VirtualBox:")
    add_bullet(doc, "забезпечує вихід віртуальної машини в глобальну мережу Інтернет для завантаження оновлень та пакетів через вбудований механізм трансляції адрес;", bold_prefix="1. Адаптер 1 (NAT): ")
    add_bullet(doc, "підключається до віртуального адаптера хоста (VirtualBox Host-Only Ethernet Adapter), створюючи виділену приватну підмережу для прямого обміну трафіком між фізичним комп'ютером і гостьовою ОС без стороннього втручання.", bold_prefix="2. Адаптер 2 (Host-Only Network): ")
    add_image_figure(doc, "fig03_vbox_nic1_nat.png", "Рис. 3. Налаштування першого мережевого адаптера віртуальної машини (режим NAT)")
    add_image_figure(doc, "fig04_vbox_nic2_hostonly.png", "Рис. 4. Налаштування другого мережевого адаптера (режим Host-Only Network)")
    add_image_figure(doc, "fig05_windows_hostonly_nic.png", "Рис. 5. Властивості інтерфейсу VirtualBox Host-Only Ethernet Adapter у Windows")

    # Step 3: Linux network & sshd
    add_header(doc, "2.3. Конфігурація мережі та розгортання служби OpenSSH в Ubuntu", 2)
    add_p(doc, "У гостьовій системі Linux задано статичну IP-адресу 192.168.56.10 із маскою 255.255.255.0 на виділеному інтерфейсі host-only. Це гарантує незмінність адреси сервера для віддалених клієнтів. Перевірено працездатність сокет-активованого демона sshd (порт 22 TCP) та стан прослуховування сокетів.")
    add_image_figure(doc, "fig06_ubuntu_network_interfaces.png", "Рис. 6. Налаштування файлу /etc/network/interfaces та статичної адреси 192.168.56.10")
    add_image_figure(doc, "fig07_sshd_service_status.png", "Рис. 7. Перевірка статусу служби ssh.service та прослуховування порту 22")

    # Step 4: sshd_config security
    add_header(doc, "2.4. Політики безпеки сервера SSH у файлі sshd_config", 2)
    add_p(doc, "Для запобігання несанкціонованому доступу та атак повного перебору (brute-force) створено конфігураційний файл /etc/ssh/sshd_config.d/lab6.conf, у якому реалізовано вимоги безпеки:")
    add_bullet(doc, "PermitRootLogin no — прямий вхід суперкористувача root заблоковано. Адміністрування здійснюється через звичайного користувача з подальшим використанням sudo або su;", bold_prefix="• Обмеження root: ")
    add_bullet(doc, "AllowUsers user remote_ws2 — доступ до сервера дозволено лише визначеним користувачам;", bold_prefix="• Білий список користувачів: ")
    add_bullet(doc, "Banner /etc/ssh/banner.txt — налаштовано обов'язкове відображення попереджувального банера Національного університету «Львівська політехніка» перед введенням облікових даних;", bold_prefix="• Попереджувальний банер: ")
    add_bullet(doc, "PasswordAuthentication yes та PubkeyAuthentication yes — увімкнено підтримку комбінованої парольної та криптографічної асиметричної автентифікації.", bold_prefix="• Методи входу: ")
    add_image_figure(doc, "fig08_sshd_config_security.png", "Рис. 8. Параметри безпеки у конфігураційному файлі sshd_config.d/lab6.conf")

    # Step 5: Network ping test
    add_header(doc, "2.5. Перевірка мережевої доступності вузла", 2)
    add_p(doc, "Перед ініціалізацією SSH-з'єднання перевірено доступність сервера за протоколом ICMP за допомогою утиліти ping 192.168.56.10, а також протестовано доступність TCP-порту 22 утилітою Test-NetConnection. Отримано 0% втрат пакетів із затримкою менше 1 мс.")
    add_image_figure(doc, "fig09_network_ping_test.png", "Рис. 9. Успішне тестування мережевого зв'язку ping та доступності порту 22")

    # Step 6: OpenSSH Client Linux
    add_header(doc, "2.6. Підключення до сервера через OpenSSH-клієнт у терміналі Linux", 2)
    add_p(doc, "Згідно із завданням 4 лабораторної роботи, виконано підключення з клієнтської машини ws2 до сервера командою ssh -l user -v 192.168.56.10. Прапорець -v активує детальний налагоджувальний вивід. Під час першого з'єднання клієнт отримав публічний ключ сервера та згенерував запит на підтвердження відбитка ключа (Host Key Fingerprint ED25519 SHA256:GQAeLagKurAME0N3AnBJ0JiLammXQ7NcVNMhiQO0k7w). Після підтвердження 'yes' ключ автоматично зафіксовано у файлі ~/.ssh/known_hosts, що захищає користувача від підміни сервера (спуфінгу).")
    add_image_figure(doc, "fig10_ssh_client_fingerprint.png", "Рис. 10. Первинний обмін протоколами та запит підтвердження відбитка відкритого ключа")
    add_image_figure(doc, "fig11_ssh_known_hosts_confirmed.png", "Рис. 11. Збереження запису хоста у відомих хостах ~/.ssh/known_hosts та введення пароля")
    add_image_figure(doc, "fig12_ssh_remote_commands.png", "Рис. 12. Виконання команд віддаленого моніторингу (uname, whoami, uptime, ps, free)")

    # Step 7: PuTTY Suite Windows
    add_header(doc, "2.7. Віддалене адміністрування з Windows за допомогою PuTTY Suite", 2)
    add_p(doc, "Для платформи Windows використано офіційний пакет PuTTY версії 0.85. У головному вікні конфігурації PuTTY створено профіль Ubuntu-SSH-Lab6, вказано IP 192.168.56.10, порт 22, а у підрозділі Connection -> Data внесено Auto-login username 'user'. Під час першого з'єднання PuTTY вивела діалогове вікно попередження PuTTY Security Alert. Після натискання кнопки [Accept] ключ хоста занесено до системного реєстру Windows HKCU\\Software\\SimonTatham\\PuTTY\\SshHostKeys.")
    add_image_figure(doc, "fig13_putty_suite_explorer.png", "Рис. 13. Склад виконуваних модулів пакета PuTTY у системному каталозі Windows")
    add_image_figure(doc, "fig14_putty_session_config.png", "Рис. 14. Створення та збереження профілю з'єднання у PuTTY Configuration")
    add_image_figure(doc, "fig15_putty_auto_login.png", "Рис. 15. Налаштування автоматичного імені користувача у Connection -> Data")
    add_image_figure(doc, "fig16_putty_security_alert.png", "Рис. 16. Попередження безпеки PuTTY Security Alert при первинному підключенні")
    add_image_figure(doc, "fig17_putty_active_session.png", "Рис. 17. Активний термінальний сеанс PuTTY з виведенням дискового простору та процесів")

    # Step 8: Keys & Pageant
    add_header(doc, "2.8. Безпарольна автентифікація за допомогою асиметричних ключів (PuTTYgen & Pageant)", 2)
    add_p(doc, "Для підвищення стійкості до компрометації облікових записів налаштовано автентифікацію за парою криптографічних ключів RSA 2048 біт. За допомогою утиліти PuTTYgen згенеровано відкритий та закритий ключі. Відкритий ключ експортовано на сервер у файл ~/.ssh/authorized_keys та встановлено обов'язкові права доступу chmod 700 на каталог ~/.ssh і chmod 600 на сам файл. Закритий ключ .ppk завантажено у резидентний агент Pageant, що дозволило виконувати миттєвий вхід на сервер без повторного введення пароля.")
    add_image_figure(doc, "fig18_puttygen_keys.png", "Рис. 18. Генерація пари ключів RSA 2048 біт в утиліті PuTTY Key Generator")
    add_image_figure(doc, "fig19_server_authorized_keys.png", "Рис. 19. Розміщення публічного ключа у файлі authorized_keys та встановлені права доступу")
    add_image_figure(doc, "fig20_pageant_agent_login.png", "Рис. 20. Завантаження ключа в агент Pageant та автоматичний вхід без введення пароля")

    # Step 9: File transfer
    add_header(doc, "2.9. Передача файлів та архівів за протоколами SCP та SFTP", 2)
    add_p(doc, "На Linux-сервері сформовано архів проєкту backup_lab6.tar.gz за допомогою утиліти tar -czf. За допомогою консольних інструментів pscp.exe та інтерактивного клієнта psftp.exe здійснено захищене завантаження та скачування архіву між Windows і Linux.")
    add_image_figure(doc, "fig21_file_transfer_scp_psftp.png", "Рис. 21. Передача архіву backup_lab6.tar.gz через pscp та інтерактивний сеанс psftp")

    # Step 10: Tunnels & Remote Desktop
    add_header(doc, "2.10. Налаштування SSH-тунелювання портів (Port Forwarding) та графічний доступ", 2)
    add_p(doc, "Для організації захищеного графічного доступу використано технологію SSH Local Port Forwarding. У PuTTY на вкладці Connection -> SSH -> Tunnels задано прокидання локального порту: Source port 3389, Destination localhost:3389. У результаті весь трафік віддаленого робочого столу RDP / X2Go інкапсулюється всередину шифрованого SSH-тунелю, що унеможливлює його перехоплення у зовнішній мережі.")
    add_image_figure(doc, "fig22_putty_ssh_tunnel_rdp.png", "Рис. 22. Конфігурація прокидання порту 3389 (L3389 localhost:3389) у PuTTY")
    add_image_figure(doc, "fig23_remote_desktop_gui.png", "Рис. 23. Підключення до робочого столу Linux через захищений тунель 127.0.0.1:3389")
    add_image_figure(doc, "fig24_server_session_audit.png", "Рис. 24. Моніторинг активних термінальних сесій та аудит сокетів sshd на сервері")

    # Control questions
    add_header(doc, "3. Відповіді на контрольні запитання", 1)
    
    q_data = [
        ("1. Чому протокол Telnet вважається небезпечним і не рекомендується до використання?",
         "Протокол Telnet передає всі дані (включно з логінами, паролями та виконуваними командами) у відкритому вигляді без будь-якого шифрування. Будь-який проміжний маршрутизатор або зловмисник у мережі може легко перехопити автентифікаційні дані за допомогою пакетних аналізаторів (Wireshark, tcpdump). Крім того, у Telnet відсутній механізм перевірки автентичності самого сервера, що робить його вразливим до атак типу Man-in-the-Middle."),
        
        ("2. У чому полягають основні переваги протоколу SSH-2 порівняно з SSH-1?",
         "SSH-2 має модульну трирівневу архітектуру (Transport, User Authentication, Connection), підтримує безпечніші алгоритми узгодження ключів (Diffie-Hellman, ECDH), надійні криптографічні шифри (AES, ChaCha20-Poly1305), алгоритми цілісності HMAC замість слабких контрольних сум CRC-32 (вразливих до атак вставки в SSH-1), а також дозволяє мультиплексувати множину термінальних та тунельованих каналів через одне TCP-з'єднання."),
        
        ("3. Яке призначення має файл ~/.ssh/known_hosts на стороні клієнта?",
         "Файл known_hosts слугує локальною базою даних відкритих ключів усіх віддалених серверів, до яких коли-небудь підключався клієнт. Під час першого підключення клієнт запам'ятовує відбиток (fingerprint) сервера. При кожному наступному сеансі клієнт звіряє наданий ключ із збереженим у known_hosts. Якщо ключ змінився, SSH виводить критичне попередження про загрозу безпеки (можлива підміна сервера або атака MITM)."),
        
        ("4. Яка різниця між парольною автентифікацією та автентифікацією за відкритим ключем (Public Key Authentication)?",
         "При парольній автентифікації користувач передає свій секретний пароль на сервер (хоч і в зашифрованому каналі), що створює ризик атаки за словником або компрометації пароля на стороні сервера. При автентифікації за відкритим ключем застосовується асиметрична криптографія: закритий ключ ніколи не покидає комп'ютер клієнта. Сервер надсилає випадкове повідомлення (челендж), яке клієнт підписує своїм закритим ключем, а сервер перевіряє підпис відкритим ключем із ~/.ssh/authorized_keys. Це забезпечує максимальний рівень захисту від компрометації."),
        
        ("5. Як працює локальне перенаправлення портів (Local Port Forwarding) у SSH?",
         "При локальному перенаправленні (наприклад, -L 3389:localhost:3389) клієнтський додаток SSH відкриває порт 3389 на комп'ютері клієнта (localhost). Будь-який локальний клієнт (наприклад, mstsc.exe), звертаючись до 127.0.0.1:3389, насправді відправляє трафік у зашифрований тунель SSH. Демон sshd на віддаленому сервері розпаковує цей трафік та перенаправляє його на кінцеву ціль (localhost:3389 на стороні сервера). Це дозволяє безпечно працювати із незахищеними протоколами через відкриті публічні мережі.")
    ]
    
    for q_title, q_ans in q_data:
        add_p(doc, q_title, bold_prefix="Запитання: ", space_after=2)
        add_p(doc, q_ans, bold_prefix="Відповідь: ", space_after=6)

    # Conclusions
    add_header(doc, "4. Висновки", 1)
    add_p(doc, "У ході виконання лабораторної роботи №6 було детально вивчено засоби та принципи безпечного віддаленого управління операційними системами Linux Ubuntu на практиці.")
    add_p(doc, "1. Налаштовано віртуальне середовище у VirtualBox із конфігурацією двох мережевих інтерфейсів: NAT для виходу у зовнішню мережу та VirtualBox Host-Only Ethernet Adapter для організації ізольованого прямого з'єднання між хостом Windows і гостьовою ОС Linux із фіксованою адресою 192.168.56.10.")
    add_p(doc, "2. Досліджено структуру та конфігураційні параметри сервера OpenSSH (sshd). У файлі sshd_config.d/lab6.conf реалізовано базові політики кібербезпеки: заборонено прямий вхід користувача root (PermitRootLogin no), активовано попереджувальний банер університету та обмежено перелік авторизованих користувачів (AllowUsers user remote_ws2).")
    add_p(doc, "3. Опановано методику віддаленого підключення через OpenSSH-клієнт у терміналі Linux із використанням налагоджувальних прапорців -l та -v. Практично досліджено процедуру первинного обміну відкритими ключами хоста, верифікацію відбитка ключа (fingerprint) та його фіксацію у файлі відомих вузлів known_hosts.")
    add_p(doc, "4. Засвоєно роботу із повнофункціональним комплектом утиліт PuTTY Suite для Windows (версія 0.85). Створено та протестовано сесійні профілі, автоматичну авторизацію (Auto-login username), обробку системного попередження PuTTY Security Alert та збереження ключів у реєстрі Windows.")
    add_p(doc, "5. Реалізовано стійку схему безпарольної автентифікації на базі пари асиметричних ключів RSA 2048 біт за допомогою PuTTY Key Generator (PuTTYgen), публічного каталогу authorized_keys та фонового агента Pageant.")
    add_p(doc, "6. Опановано захищену передачу файлових архівів за протоколами SCP та SFTP (утиліти pscp, psftp), а також налаштовано технологію SSH Port Forwarding для інкапсуляції графічного трафіку віддаленого робочого столу (порт 3389).")

    report_path = os.path.join(REPORTS_DIR, "Лемішка_ПП_14_Лаб_6.docx")
    doc.save(report_path)
    print("Saved Main Report DOCX:", report_path)
    return report_path

# ==========================================================
# BUILD DOCUMENT 2: RUNBOOK (Покрокова_інструкція_натискань_Лаб6.docx)
# ==========================================================
def build_runbook_doc():
    print("Building Step-by-Step Instruction Runbook: Покрокова_інструкція_натискань_Лаб6.docx...")
    doc = Document()
    setup_document_styles(doc)
    add_title_page(doc, doc_type="ПОКРОКОВА ІНСТРУКЦІЯ", lab_num="6", theme="«Покроковий посібник виконання та натискань дій: Віддалений доступ в Linux (SSH, PuTTY, SFTP, Tunnels)»")

    add_header(doc, "Покрокова інструкція виконання дій до лабораторної роботи №6", 1)
    add_p(doc, "Цей посібник містить вичерпний опис усіх 24 послідовних кроків, натискань кнопок у графічних діалогових вікнах (VirtualBox, PuTTY, PuTTYgen, Pageant, mstsc) та команд термінала для бездоганного проходження та захисту лабораторної роботи.")

    steps = [
        ("Крок 1. Створення віртуальної машини у VirtualBox",
         "1. Відкрити вікно Oracle VM VirtualBox.\n2. Натиснути кнопку «Створити» (Create / New) на панелі інструментів.\n3. У полі «Name» ввести назву машини: Ubuntu_Lab6.\n4. Перевірити шлях у полі «Folder»: C:\\Users\\user\\VirtualBox VMs.\n5. Вибрати Type: «Linux», Version: «Ubuntu (64-bit)».\n6. Встановити галочку «Skip Unattended Installation» -> Натиснути кнопку «Next >».",
         "fig01_vbox_create_wizard.png", "Рис. 1.1. Заповнення полів майстра створення віртуальної машини Ubuntu_Lab6"),

        ("Крок 2. Виділення ресурсів RAM та накопичувача VDI",
         "1. У розділі «Hardware» встановити повзунок Base Memory на позначку 2048 MB (2 GB RAM).\n2. Встановити кількість процесорів: 2 CPUs.\n3. У розділі «Virtual Harddisk» вибрати «Create a Virtual Hard Disk Now».\n4. Задати розмір накопичувача: 20.00 GB.\n5. Тип накопичувача — VDI (VirtualBox Disk Image), динамічне виділення.\n6. Натиснути кнопку «Finish».",
         "fig02_vbox_ram_and_disk.png", "Рис. 1.2. Виділення пам'яті 2048 MB та диска 20 GB у майстрі"),

        ("Крок 3. Налаштування першого мережевого адаптера (NAT)",
         "1. Вибрати віртуальну машину Ubuntu_Lab6 у списку VirtualBox -> натиснути кнопку «Налаштування» (Settings).\n2. Перейти у лівому меню до розділу «Network» (Мережа).\n3. Відкрити вкладку «Adapter 1».\n4. Переконатися, що прапорець «Enable Network Adapter» увімкнено.\n5. У списку «Attached to:» вибрати «NAT».\n6. Розгорнути вкладку «Advanced» і перевірити тип адаптера Intel PRO/1000 MT Desktop.",
         "fig03_vbox_nic1_nat.png", "Рис. 1.3. Конфігурація Адаптера 1 у режимі NAT для виходу в Інтернет"),

        ("Крок 4. Налаштування другого адаптера (Host-Only Network)",
         "1. У вікні налаштувань мережі перейти на вкладку «Adapter 2».\n2. Встановити прапорець «Enable Network Adapter».\n3. У випадаючому списку «Attached to:» вибрати «Host-only Adapter» (Віртуальний адаптер хоста).\n4. У списку «Name» вибрати «VirtualBox Host-Only Ethernet Adapter».\n5. У вкладці «Advanced» вибрати Promiscuous Mode: «Allow All».\n6. Натиснути кнопку «OK» для збереження змін.",
         "fig04_vbox_nic2_hostonly.png", "Рис. 1.4. Підключення другого адаптера до виділеної мережі Host-Only"),

        ("Крок 5. Перевірка хостового адаптера у Панелі керування Windows",
         "1. Відкрити Панель керування Windows -> «Центр керування мережами» -> «Змінити параметри адаптера» (ncpa.cpl).\n2. Знайти підключення «VirtualBox Host-Only Network».\n3. Натиснути правою кнопкою миші -> «Властивості» (Properties).\n4. Переконатися, що активовано VirtualBox NDIS6 Bridged Networking Driver та TCP/IPv4 із адресою 192.168.56.1.\n5. Натиснути кнопку «OK».",
         "fig05_windows_hostonly_nic.png", "Рис. 1.5. Властивості інтерфейсу VirtualBox Host-Only у Windows 11"),

        ("Крок 6. Налаштування статичного IP у Linux (/etc/network/interfaces)",
         "1. Запустити термінал Ubuntu.\n2. Відкрити файл конфігурації мережі: sudo nano /etc/network/interfaces.\n3. Внести параметри інтерфейсу enp0s8:\n   auto enp0s8\n   iface enp0s8 inet static\n   address 192.168.56.10\n   netmask 255.255.255.0\n4. Зберегти файл (Ctrl+O, Enter, Ctrl+X) та перевірити активну адресу командою ip addr show.",
         "fig06_ubuntu_network_interfaces.png", "Рис. 1.6. Файл /etc/network/interfaces зі статичною адресою 192.168.56.10"),

        ("Крок 7. Перевірка статусу служби OpenSSH (sshd)",
         "1. Виконати команду: sudo systemctl status ssh.\n2. Перевірити наявність зеленого статусу «active (running)».\n3. Перевірити сокети командою: sudo ss -tulpn | grep :22.\n4. Переконатися, що порт 22 відкрито для прослуховування на адресах 0.0.0.0:22 та [::]:22.",
         "fig07_sshd_service_status.png", "Рис. 1.7. Виведення статусу служби ssh.service та прослуховування 22 порту"),

        ("Крок 8. Налаштування параметрів безпеки у файлі sshd_config",
         "1. Відкрити каталог налаштувань: cd /etc/ssh/sshd_config.d/.\n2. Створити конфігураційний файл lab6.conf: sudo nano lab6.conf.\n3. Внести параметри:\n   Port 22\n   Banner /etc/ssh/banner.txt\n   PermitRootLogin no\n   PasswordAuthentication yes\n   PubkeyAuthentication yes\n   AllowUsers user remote_ws2\n4. Перевірити синтаксис: sudo sshd -t -> перезавантажити службу: sudo systemctl restart ssh.",
         "fig08_sshd_config_security.png", "Рис. 1.8. Застосування параметрів безпеки sshd_config.d/lab6.conf"),

        ("Крок 9. Перевірка мережевого зв'язку ping та порту 22",
         "1. На хост-машині Windows відкрити термінал PowerShell.\n2. Виконати команду: ping 192.168.56.10.\n3. Переконатися, що надійшло 4 відповіді (0% втрат, затримка < 1 мс).\n4. Протестувати порт 22 командою: Test-NetConnection -ComputerName 192.168.56.10 -Port 22.\n5. Перевірити значення TcpTestSucceeded : True.",
         "fig09_network_ping_test.png", "Рис. 1.9. Успішне проходження пакетів ICMP та перевірка відкритого порту 22"),

        ("Крок 10. Підключення через OpenSSH-клієнт у терміналі Linux",
         "1. На сусідній клієнтській машині Linux виконати команду:\n   ssh -l user -v 192.168.56.10\n2. Спостерігати за процесом налагодження OpenSSH (обмін версіями протоколу SSH-2.0).\n3. Отримати запит верифікації ключа сервера (The authenticity of host cannot be established).\n4. Перевірити відповідність відбитка SHA256:GQAeLagKurAME0N3AnBJ0JiLammXQ7NcVNMhiQO0k7w.",
         "fig10_ssh_client_fingerprint.png", "Рис. 1.10. Відображення відбитка ключа хоста при першому підключенні"),

        ("Крок 11. Підтвердження відбитка та збереження у known_hosts",
         "1. У відповідь на запит впевненості ввести з клавіатури: yes.\n2. Отримати повідомлення: Warning: Permanently added '192.168.56.10' to the list of known hosts.\n3. Ввести пароль користувача (12345).\n4. Перевірити створений запис у файлі відомих хостів командою: tail -n 1 ~/.ssh/known_hosts.",
         "fig11_ssh_known_hosts_confirmed.png", "Рис. 1.11. Фіксація запису хоста у known_hosts та введення пароля"),

        ("Крок 12. Виконання команд віддаленого моніторингу на сервері",
         "1. Переглянути університетський банер вітання Львівської політехніки.\n2. Виконати системні команди моніторингу:\n   uname -a (версія ядра);\n   whoami (поточний користувач);\n   uptime (час безперервної роботи);\n   ps -ef | grep pts (активні процеси термінала);\n   free -m (обсяг пам'яті).",
         "fig12_ssh_remote_commands.png", "Рис. 1.12. Інтерактивне виконання діагностичних команд у сесії SSH"),

        ("Крок 13. Огляд дистрибутива PuTTY Suite у Windows",
         "1. Відкрити Провідник Windows (File Explorer).\n2. Перейти до каталогу C:\\Program Files\\PuTTY.\n3. Ознайомитися з файлами дистрибутива: putty.exe (клієнт), puttygen.exe (генератор ключів), pageant.exe (агент ключів), pscp.exe (копіювання SCP), psftp.exe (клієнт SFTP), plink.exe (CLI-інтерфейс).",
         "fig13_putty_suite_explorer.png", "Рис. 1.13. Виконувані файли пакета PuTTY 0.85 у Провіднику Windows"),

        ("Крок 14. Створення профілю сесії у PuTTY Configuration",
         "1. Запустити putty.exe.\n2. У полі «Host Name (or IP address)» ввести: 192.168.56.10.\n3. Переконатися, що встановлено Port: 22 та Connection type: SSH.\n4. У полі «Saved Sessions» ввести ім'я сесії: Ubuntu-SSH-Lab6.\n5. Натиснути кнопку «Save» для збереження налаштувань у системному профілі.",
         "fig14_putty_session_config.png", "Рис. 1.14. Збереження профілю Ubuntu-SSH-Lab6 у PuTTY"),

        ("Крок 15. Налаштування автоматичного логіна у Connection -> Data",
         "1. У лівому дереві категорій PuTTY розгорнути вузол «Connection» -> перейти до підрозділу «Data».\n2. У полі «Auto-login username» ввести: user.\n3. Переконатися, що Terminal-type string встановлено як xterm-256color.\n4. Повернутися на вкладку «Session» та повторно натиснути кнопку «Save».\n5. Натиснути кнопку «Open» для підключення.",
         "fig15_putty_auto_login.png", "Рис. 1.15. Встановлення автоматичного імені входу user"),

        ("Крок 16. Обробка попередження безпеки PuTTY Security Alert",
         "1. Під час першого з'єднання з'явиться вікно «PuTTY Security Alert».\n2. Ознайомитися із текстом про відсутність ключа в реєстрі Windows.\n3. Перевірити збіг відбитка ключа із сервером: ssh-ed25519 255 SHA256:GQAeLagKurAME0N3AnBJ0JiLammXQ7NcVNMhiQO0k7w.\n4. Натиснути кнопку «Accept» для додавання ключа до реєстру Windows.",
         "fig16_putty_security_alert.png", "Рис. 1.16. Натискання кнопки [Accept] у PuTTY Security Alert"),

        ("Крок 17. Сеанс віддаленого адміністрування у вікні PuTTY",
         "1. У вікні термінала PuTTY ввести пароль користувача.\n2. Переглянути університетський банер вітання.\n3. Виконати команди моніторингу: df -h (стан дискових розділів), last -n 4 (журнал входів).\n4. Перевірити коректність завершення сеансу командою logout або exit.",
         "fig17_putty_active_session.png", "Рис. 1.17. Робоче вікно термінала PuTTY після успішного входу"),

        ("Крок 18. Генерація пари ключів у PuTTY Key Generator (PuTTYgen)",
         "1. Запустити утиліту puttygen.exe.\n2. У розділі «Parameters» вибрати тип ключа: RSA, довжина: 2048 біт.\n3. Натиснути кнопку «Generate» та переміщати мишу у вікні для генерації випадкової ентропії.\n4. У полі «Key comment» ввести: rsa-key-lemishka-2026.\n5. Ввести та підтвердити парольну фразу (Key passphrase).\n6. Натиснути «Save private key» і зберегти файл private_lemishka.ppk.\n7. Скопіювати текст публічного ключа із верхнього текстового поля.",
         "fig18_puttygen_keys.png", "Рис. 1.18. Генерація та збереження асиметричної пари ключів у PuTTYgen"),

        ("Крок 19. Розміщення публічного ключа у файлі authorized_keys",
         "1. На сервері Ubuntu перейти до домашнього каталогу користувача.\n2. Створити каталог ~/.ssh і файл authorized_keys: mkdir -p ~/.ssh && nano ~/.ssh/authorized_keys.\n3. Вставити скопійований публічний ключ.\n4. Встановити суворі права доступу:\n   chmod 700 ~/.ssh\n   chmod 600 ~/.ssh/authorized_keys\n5. Перевірити права командою ls -la ~/.ssh.",
         "fig19_server_authorized_keys.png", "Рис. 1.19. Налаштування прав доступу до файлу authorized_keys"),

        ("Крок 20. Завантаження ключа в Pageant та безпарольний вхід",
         "1. Запустити агент ключів pageant.exe.\n2. У системному треї двічі клацнути на значок Pageant -> натиснути кнопку «Add Key».\n3. Вибрати збережений файл private_lemishka.ppk -> ввести парольну фразу.\n4. Відкрити сесію PuTTY: переконатися, що вхід відбувається миттєво без запиту пароля (Authenticating with public key from agent).",
         "fig20_pageant_agent_login.png", "Рис. 1.20. Вікно Pageant Key List та автоматична автентифікація за ключем"),

        ("Крок 21. Передача файлів та архівів через pscp та psftp",
         "1. На сервері запакувати проєкт: tar -czf backup_lab6.tar.gz lab6_project/.\n2. На Windows передати файл через pscp: pscp.exe -batch -pw 12345 backup_lab6.tar.gz user@192.168.56.10:/home/user/incoming/.\n3. Запустити psftp.exe user@192.168.56.10.\n4. Виконати команди: ls, get incoming/backup_lab6.tar.gz local_copy.tar.gz, quit.",
         "fig21_file_transfer_scp_psftp.png", "Рис. 1.21. Передача архівів та робота у сесії клієнта psftp"),

        ("Крок 22. Налаштування SSH-тунелювання портів (Port Forwarding)",
         "1. Відкрити PuTTY Configuration -> перейти до Connection -> SSH -> Tunnels.\n2. У полі «Source port» ввести: 3389.\n3. У полі «Destination» ввести: localhost:3389.\n4. Вибрати радіокнопку «Local» -> натиснути кнопку «Add».\n5. Перевірити доданий запис у списку Forwarded ports: L3389 localhost:3389.\n6. Зберегти профіль сесії.",
         "fig22_putty_ssh_tunnel_rdp.png", "Рис. 1.22. Додавання правила тунелювання порту 3389 у PuTTY"),

        ("Крок 23. Підключення до робочого столу через тунель (127.0.0.1:3389)",
         "1. Переконатися, що SSH-сесія із тунелем активна.\n2. Натиснути Win+R -> ввести mstsc.exe -> натиснути Enter.\n3. У полі «Комп'ютер» (Computer) ввести: 127.0.0.1:3389.\n4. У полі «Користувач» ввести: user.\n5. Натиснути кнопку «Підключити» (Connect) -> ввести облікові дані для входу у графічне середовище Linux через захищений тунель.",
         "fig23_remote_desktop_gui.png", "Рис. 1.23. Діалогове вікно підключення mstsc до тунельованої адреси 127.0.0.1:3389"),

        ("Крок 24. Моніторинг активних підключень та аудит безпеки",
         "1. На сервері Ubuntu виконати перевірку активних сесій: w та who -u.\n2. Перевірити активні з'єднання: ss -tna '( dport = :22 or sport = :22 )'.\n3. Переглянути системний журнал автентифікації: journalctl -u ssh -n 10.\n4. Переконатися у відсутності підозрілих невдалих спроб входу.",
         "fig24_server_session_audit.png", "Рис. 1.24. Моніторинг активних сокетів та журналу автентифікації на сервері")
    ]

    for title, desc, img, cap in steps:
        add_header(doc, title, 2)
        add_p(doc, desc)
        add_image_figure(doc, img, cap)

    runbook_path = os.path.join(REPORTS_DIR, "Покрокова_інструкція_натискань_Лаб6.docx")
    doc.save(runbook_path)
    print("Saved Runbook DOCX:", runbook_path)
    return runbook_path

# ==========================================================
# EXPORT TO PDF VIA WORD COM
# ==========================================================
def convert_to_pdf(docx_path, pdf_path):
    print(f"Exporting PDF via Word COM: {docx_path} -> {pdf_path}...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(docx_path, False, True)
        doc.SaveAs(pdf_path, 17) # 17 = wdFormatPDF
        doc.Close(False)
        print("Export successful!")
    except Exception as e:
        print("Error during PDF export:", e)
    finally:
        word.Quit()
        
    if os.path.exists(pdf_path):
        sz = os.path.getsize(pdf_path) / (1024 * 1024)
        print(f"PDF created: {pdf_path} ({sz:.2f} MB)")
        if sz > 5.0:
            print("WARNING: File size exceeds 5MB!")
        else:
            print("PASSED: File size is strictly under 5MB.")

def main():
    print("==========================================================")
    print("Generating Complete Laboratory Work 6 Documentation Suite")
    print("==========================================================")
    
    # Build both DOCX files
    rep_docx = build_report_doc()
    run_docx = build_runbook_doc()
    
    # Target PDF paths in reports/
    rep_pdf = os.path.join(REPORTS_DIR, "Лемішка_ПП_14_Лаб_6.pdf")
    run_pdf = os.path.join(REPORTS_DIR, "Покрокова_інструкція_натискань_Лаб6.pdf")
    
    # Export to PDF
    convert_to_pdf(rep_docx, rep_pdf)
    convert_to_pdf(run_docx, run_pdf)
    
    # Copy to Lab6/ root for easy student access
    root_rep_docx = os.path.join(LAB_DIR, "Лемішка_ПП_14_Лаб_6.docx")
    root_rep_pdf = os.path.join(LAB_DIR, "Лемішка_ПП_14_Лаб_6.pdf")
    root_run_docx = os.path.join(LAB_DIR, "Покрокова_інструкція_натискань_Лаб6.docx")
    root_run_pdf = os.path.join(LAB_DIR, "Покрокова_інструкція_натискань_Лаб6.pdf")
    
    shutil.copyfile(rep_docx, root_rep_docx)
    shutil.copyfile(rep_pdf, root_rep_pdf)
    shutil.copyfile(run_docx, root_run_docx)
    shutil.copyfile(run_pdf, root_run_pdf)
    
    print("\nAll files successfully created and copied to root of Lab6:")
    for f in [root_rep_docx, root_rep_pdf, root_run_docx, root_run_pdf]:
        sz = os.path.getsize(f) / (1024 * 1024)
        print(f" - {os.path.basename(f)}: {sz:.2f} MB")

if __name__ == '__main__':
    main()
