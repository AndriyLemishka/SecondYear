# -*- coding: utf-8 -*-
"""
Генератор пам'ятки / підготовки до захисту лабораторної роботи №6
Проєкт: ChargeUA
Студент: Лемішка Андрій Андрійович, ПП-24
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = r"D:\Studing\SecondYear\FirstSemester\web\lab6"
DOCX_OUT = os.path.join(BASE_DIR, "Лемішка_ПП-24_Л6_Підготовка_до_захисту.docx")
PDF_OUT = os.path.join(BASE_DIR, "Лемішка_ПП-24_Л6_Підготовка_до_захисту.pdf")

COLOR_PRIMARY = RGBColor(0, 94, 201)     # #005EC9 (ChargeUA Blue)
COLOR_DARK = RGBColor(15, 23, 42)        # #0F172A
COLOR_MUTED = RGBColor(71, 85, 105)      # #475569
COLOR_SUCCESS = RGBColor(16, 185, 129)   # #10B981
COLOR_BORDER = "CBD5E1"
COLOR_BG_LIGHT = "F8FAFC"
COLOR_BG_CARD = "F1F5F9"
COLOR_BG_PRIMARY_TINT = "EFF6FF"

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_bg(cell, hex_color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='005EC9')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for b_name in ['top', 'left', 'bottom', 'right']:
        if b_name in kwargs:
            edge = OxmlElement(f'w:{b_name}')
            for key, val in kwargs[b_name].items():
                edge.set(qn(f'w:{key}'), str(val))
            tcBorders.append(edge)
        else:
            edge = OxmlElement(f'w:{b_name}')
            edge.set(qn('w:val'), 'none')
            tcBorders.append(edge)
    tcPr.append(tcBorders)

def build_defense_doc():
    doc = docx.Document()
    
    # Page setup (A4, 1.5 cm margins for dense, beautiful 2-3 page cheat sheet)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = COLOR_DARK

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=2, space_after=2, line_spacing=1.15):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            r.font.name = 'Calibri'
        return p

    def add_h1(title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.name = 'Calibri'
        r.font.size = Pt(13)
        r.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        
        # Add thin bottom border / line via XML
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '12')
        bottom.set(qn('w:space'), '4')
        bottom.set(qn('w:color'), '005EC9')
        pBdr.append(bottom)
        pPr.append(pBdr)
        return p

    def add_callout(text, prefix="💡 ПОРАДА: "):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(7.07)
        set_cell_bg(cell, COLOR_BG_PRIMARY_TINT)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        set_cell_border(cell, 
            left=dict(sz=24, val='single', color='005EC9'),
            top=dict(sz=4, val='single', color='E2E8F0'),
            bottom=dict(sz=4, val='single', color='E2E8F0'),
            right=dict(sz=4, val='single', color='E2E8F0')
        )
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r_pre = p.add_run(prefix)
        r_pre.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY
        r_txt = p.add_run(text)
        r_txt.font.size = Pt(9.5)

    def add_qa_card(q_num, question, answer, code_tip=None):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(7.07)
        set_cell_bg(cell, COLOR_BG_LIGHT)
        set_cell_margins(cell, top=100, bottom=120, left=140, right=140)
        set_cell_border(cell,
            left=dict(sz=16, val='single', color='005EC9'),
            top=dict(sz=4, val='single', color='E2E8F0'),
            bottom=dict(sz=4, val='single', color='E2E8F0'),
            right=dict(sz=4, val='single', color='E2E8F0')
        )
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        r_q = p.add_run(f"❓ {q_num}. {question}\n")
        r_q.bold = True
        r_q.font.size = Pt(10)
        r_q.font.color.rgb = COLOR_DARK

        p_a = cell.add_paragraph()
        p_a.paragraph_format.space_before = Pt(2)
        p_a.paragraph_format.space_after = Pt(2)
        p_a.paragraph_format.line_spacing = 1.15
        r_ans = p_a.add_run("🎯 Відповідь: ")
        r_ans.bold = True
        r_ans.font.size = Pt(9.5)
        r_ans.font.color.rgb = COLOR_PRIMARY
        r_text = p_a.add_run(answer)
        r_text.font.size = Pt(9.5)

        if code_tip:
            p_c = cell.add_paragraph()
            p_c.paragraph_format.space_before = Pt(3)
            p_c.paragraph_format.space_after = Pt(2)
            r_c_pre = p_c.add_run("💻 Що показати / ключовий код: ")
            r_c_pre.bold = True
            r_c_pre.font.size = Pt(9)
            r_c_pre.font.color.rgb = RGBColor(100, 116, 139)
            r_c_val = p_c.add_run(code_tip)
            r_c_val.font.size = Pt(9)
            r_c_val.font.name = 'Consolas'
            r_c_val.font.color.rgb = RGBColor(30, 41, 59)

        # Empty spacing
        p_end = doc.add_paragraph()
        p_end.paragraph_format.space_before = Pt(2)
        p_end.paragraph_format.space_after = Pt(2)

    # =========================================================================
    # HEADER / TITLE BANNER
    # =========================================================================
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_cell = header_table.cell(0, 0)
    h_cell.width = Inches(7.07)
    set_cell_bg(h_cell, "0F172A") # Dark Navy
    set_cell_margins(h_cell, top=140, bottom=140, left=180, right=180)
    
    p_top = h_cell.paragraphs[0]
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_top = p_top.add_run("НУ «ЛЬВІВСЬКА ПОЛІТЕХНІКА» · КАФЕДРА СИСТЕМ ШТУЧНОГО ІНТЕЛЕКТУ\n")
    r_top.font.size = Pt(8.5)
    r_top.font.color.rgb = RGBColor(148, 163, 184)
    r_top.bold = True

    r_main = p_top.add_run("ПАМ'ЯТКА ТА ШПАРГАЛКА ДО ЗАХИСТУ ЛАБОРАТОРНОЇ РОБОТИ №6\n")
    r_main.font.size = Pt(14)
    r_main.font.color.rgb = RGBColor(255, 255, 255)
    r_main.bold = True

    r_sub = p_top.add_run("Тема: Адаптивна компонентна верстка інтерфейсу вебсайту за допомогою Tailwind CSS\n")
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(56, 189, 248) # light blue

    r_meta = p_top.add_run("Студент: Лемішка Андрій Андрійович (група ПП-24) · Проєкт: ChargeUA\n"
                           "Живий деплой: https://mellifluous-eclair-8fe561.netlify.app/ · GitHub: AndriyLemishka/SecondYear")
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = RGBColor(203, 213, 225)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # =========================================================================
    # РОЗДІЛ 1: 90-СЕКУНДНИЙ ВСТУПНИЙ СПІЧ (ELEVATOR PITCH)
    # =========================================================================
    add_h1("1. ВСТУПНИЙ СПІЧ НА ЗАХИСТІ (ЯК ПРЕЗЕНТУВАТИ ЗА 90 СЕКУНД)")
    add_callout(
        "«Доброго дня! У Лабораторній роботі №6 я виконав рефакторинг та розробку адаптивних компонентів "
        "для національного агрегатора зарядних станцій ChargeUA з використанням Tailwind CSS.\n"
        "1. Основна концепція: ChargeUA — це інформаційний хаб і калькулятор (ми не проводимо транзакції і не бронюємо, "
        "а допомагаємо водію електромобіля знайти швидкісний порт і розрахувати вартість сесії).\n"
        "2. Архітектура: застосовано гібридний підхід. Ми зберегли дизайн-токени (CSS Custom Properties у :root) "
        "і поєднали їх з utility-first класами Tailwind CSS (container, flex, grid, rounded-xl, sm:flex-row).\n"
        "3. Mobile-First: реалізовано адаптивну навігацію з доступним гамбургер-меню (WCAG 2.2, керування через клавіатуру, Escape, блокування фонового скролу), "
        "гнучкі сітки для 375px (мобільні), 768px (планшети) та 1440px (десктоп).\n"
        "4. Продуктивність та оптимізація: карту Leaflet та калькулятор переведено на ліниве завантаження через IntersectionObserver, "
        "а стилі Tailwind скомпільовано в tailwind.min.css. Завдяки цьому в аудиті Google Lighthouse ми досягли: "
        "Best Practices — 100%, SEO — 100%, Accessibility — 97%, Performance — 87% з ідеальним Total Blocking Time 0 мс та CLS 0.\n"
        "Проєкт повністю валідний за стандартом W3C (0 помилок) і опублікований на Netlify. Готовий продемонструвати функціонал і код!»",
        prefix="🎤 ВАШ ГОТОВИЙ СПІЧ: "
    )

    # =========================================================================
    # РОЗДІЛ 2: ТОП-8 ПИТАНЬ ТА ІДЕАЛЬНИХ ВІДПОВІДЕЙ
    # =========================================================================
    add_h1("2. ТОП-8 ПИТАНЬ ВИКЛАДАЧА ТА ВПЕВНЕНІ ВІДПОВІДІ (БЛІЦ-ОПИТУВАННЯ)")

    add_qa_card(
        1,
        "Чим Tailwind CSS відрізняється від компонентних фреймворків на зразок Bootstrap?",
        "Bootstrap дає готові монолітні компоненти (.btn, .card, .modal) з наперед визначеним дизайном, через що сайти виглядають однотипно, "
        "а перевизначення стилів потребує громіздкого CSS. Tailwind CSS — це utility-first фреймворк, що надає атомарні низькорівневі утиліти "
        "(flex, p-4, bg-blue-600, rounded-xl). Це дає абсолютну свободу створення унікального дизайну без написання кастомного CSS-коду. "
        "Крім того, завдяки JIT-компіляції у фінальний бандл потрапляють лише реально використані класи, тому результуючий CSS важить лічені кілобайти.",
        "index.html: class='card p-4 rounded-xl shadow-sm bg-white border border-slate-200'"
    )

    add_qa_card(
        2,
        "Як у проєкті поєднано власний style.css та Tailwind CSS?",
        "Ми використали гібридну архітектуру. У style.css збережено глобальні дизайн-токени бренду ChargeUA (:root з фірмовими кольорами, "
        "розмірами шрифтів clamp(), відступами та базовим скиданням CSS Reset). Натомість для сіток, карток станцій, адаптивних лейаутів "
        "та інтерактивних станів використано Tailwind CSS. Це забезпечує стабільність дизайн-системи та швидкість розробки компонентів.",
        "style.css (:root { --color-primary: #005EC9; }) + tailwind.min.css (утиліти)"
    )

    add_qa_card(
        3,
        "Як реалізовано принцип Mobile-First у Tailwind та чистому CSS?",
        "Принцип Mobile-First означає, що базові стилі (без префіксів) пишуться для найменших мобільних екранів (320–375px), "
        "а для ширших екранів стилі послідовно розширюються через медіа-запити min-width. У Tailwind це працює автоматично: "
        "класи без префікса (наприклад, flex flex-col) діють на мобільних, а префікси sm: (від 640px) та md: (від 768px) активуються "
        "лише при розширенні вікна браузера. Усі сенсорні зони (кнопки, чипси) мають мінімум 44×44px для зручного натискання пальцем.",
        "class='flex flex-col sm:flex-row md:items-end gap-3' (1 колонка на мобільних -> горизонтально від 640px)"
    )

    add_qa_card(
        4,
        "Як працює гамбургер-меню з точки зору веб-доступності (WCAG 2.2 AA)?",
        "Модуль js/nav.js забезпечує повну доступність для клавіатури та скрінрідерів: кнопка має семантичний тег <button>, "
        "атрибути aria-label='Відкрити меню навігації', aria-controls='site-nav' та динамічний aria-expanded='true/false'. "
        "При відкритті меню на мобільному скрол сторінки блокується (overflow: hidden), фокус залишається доступним, "
        "а закриття реалізовано трьома способами: кліком на кнопку-хрестик, натисканням клавіші Escape, або кліком на будь-яке посилання навігації.",
        "js/nav.js: burgerBtn.setAttribute('aria-expanded', isExpanded); document.addEventListener('keydown', (e) => e.key === 'Escape' && closeNav())"
    )

    add_qa_card(
        5,
        "Як влаштований інтерактивний калькулятор зарядки ChargeUA?",
        "Калькулятор реалізовано на чистому JavaScript (js/calculator.js). Він містить базу даних понад 45 моделей електромобілів "
        "(Tesla, VW, Nissan, BMW, BYD) з реальними параметрами: ємність батареї, максимальна потужність DC, тип порту та розхід енергії. "
        "Користувач обирає авто (або вводить кВт·год вручну) і рухає інтерактивний повзунок від 20% до 80%. "
        "Скрипт динамічно розраховує: додані кВт·год, вартість за тарифом 14.12 ₴/кВт·год, час зарядки на станції 120 кВт та економію проти бензину А-95.",
        "js/calculator.js: sessionCost = addedKwh * car.tariff; savings = benzinCost - sessionCost"
    )

    add_qa_card(
        6,
        "Як реалізовано ліниве завантаження карти (Leaflet) та чому це важливо?",
        "Інтерактивна карта Leaflet та тайли Esri ArcGIS важать понад 280 КБ і вимагають десятків мережевих HTTP-запитів. "
        "Щоб карта не гальмувала перший рендеринг сторінки, ми використали IntersectionObserver: сам скрипт leaflet.js і ініціалізація карти "
        "запускаються лише тоді, коли блок #map реально перетинає вікно перегляду (rootMargin: '0px'). "
        "Під час початкового завантаження сторінки час навантаження процесора (CPU) дорівнює 0 мс.",
        "js/map.js: observer = new IntersectionObserver((entries) => { if (entries[0].isIntersecting) loadLeaflet(initMap); })"
    )

    add_qa_card(
        7,
        "Які результати аудитів W3C та Google Lighthouse?",
        "1. W3C Nu HTML Checker: перевірка сторінки index.html показала 0 помилок і 0 попереджень (повна валідність HTML5).\n"
        "2. Google Lighthouse: на опублікованому на Netlify сайті досягнуто: Best Practices — 100/100, SEO — 100/100, "
        "Accessibility — 97/100 (зразковий рівень WCAG) та Performance — 87/100 з Total Blocking Time (TBT) рівно 0 мс та зсувом макета (CLS) рівно 0.",
        "Звіт Lighthouse зафіксовано на Рис. 15 у розділі 4 звіту."
    )

    add_qa_card(
        8,
        "Чому в проєкті ChargeUA немає бронювання та оплати онлайн?",
        "Згідно з технічним завданням, ChargeUA є незалежним національним агрегатором станцій різних операторів (YASNO, AutoEnterprise, TOKA, EcoFactor). "
        "Агрегатор надає інформаційні послуги: пошук вільних портів у реальному часі, фільтрацію за конекторами, карту покриття та розрахунок вартості. "
        "Він свідомо не проводить транзакцій і не бере комісій, направляючи водія до відповідного застосунку оператора.",
        "Розділ 'Юридичні аспекти агрегатора' у звіті."
    )

    # =========================================================================
    # РОЗДІЛ 3: ЖИВА ДЕМОНСТРАЦІЯ (ЧЕКЛИСТ ДЛЯ ПОКАЗУ)
    # =========================================================================
    add_h1("3. ПОКРОКОВИЙ ПЛАН ДЕМОНСТРАЦІЇ (ЩО ПОКАЗАТИ ВИКЛАДАЧУ ЗА 60 СЕКУНД)")

    demo_steps = [
        ("Крок 1. Відкрити живий сайт у браузері", "https://mellifluous-eclair-8fe561.netlify.app/ — показати швидке завантаження, адаптивний хедер, логотип та чіткий геро-банер."),
        ("Крок 2. Продемонструвати калькулятор зарядки", "Натиснути на швидкий пресет авто (наприклад, Tesla або VW ID.4), порухати повзунок відсотка заряду — показати, як миттєво змінюється ціна, час сесії та економія."),
        ("Крок 3. Натиснути 'Знайти станцію'", "Клік на кнопку плавно скролить сторінку до карти та автоматично активує фільтр конектора під обране авто (наприклад, CCS 2)."),
        ("Крок 4. Показати інтерактивну карту", "Показати маркери станцій (Львів, Київ, траси), клікнути на маркер — відкривається поп-ап з живою кількістю вільних портів і тарифом."),
        ("Крок 5. Фільтрація каталогу станцій", "Клікнути на фільтр міст (Львів, Київ, Одеса) або кнопку 'Тільки вільні' — показати миттєве оновлення карток та кнопку 'Показати ще'."),
        ("Крок 6. DevTools — режим смартфона (375px)", "Натиснути F12 -> іконка смартфона (375px): показати роботу гамбургер-меню (відкриття/закриття/Escape) та 1-колонковий адаптивний лейаут."),
        ("Крок 7. Показати код у VS Code", "Відкрити index.html (показати класи Tailwind, семантичні теги header/main/section/footer), style.css (:root токени) та скрипти у js/.")
    ]

    table_demo = doc.add_table(rows=len(demo_steps) + 1, cols=2)
    table_demo.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_demo.autofit = False

    # Header row
    hdr_cells = table_demo.rows[0].cells
    hdr_cells[0].width = Inches(2.3)
    hdr_cells[1].width = Inches(4.77)
    hdr_cells[0].text = "Дія / Етап показу"
    hdr_cells[1].text = "Що саме коментувати викладачу"
    for c in hdr_cells:
        set_cell_bg(c, "0F172A")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.runs[0].bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9.5)

    for i, (action, comment) in enumerate(demo_steps):
        row_cells = table_demo.rows[i + 1].cells
        row_cells[0].width = Inches(2.3)
        row_cells[1].width = Inches(4.77)
        row_cells[0].text = action
        row_cells[1].text = comment
        bg = COLOR_BG_LIGHT if i % 2 == 0 else "FFFFFF"
        for c_idx, c in enumerate(row_cells):
            set_cell_bg(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)
            set_cell_border(c, 
                bottom=dict(sz=4, val='single', color=COLOR_BORDER),
                top=dict(sz=4, val='single', color=COLOR_BORDER)
            )
            p = c.paragraphs[0]
            p.runs[0].font.size = Pt(9)
            if c_idx == 0:
                p.runs[0].bold = True
                p.runs[0].font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # =========================================================================
    # РОЗДІЛ 4: ПІДСТУПНІ ЗАПИТАННЯ (ТРАП-ПИТАННЯ)
    # =========================================================================
    add_h1("4. ПІДСТУПНІ ЗАПИТАННЯ ВИКЛАДАЧА («TRICKY QUESTIONS»)")

    add_qa_card(
        "Т1",
        "«А якщо у клієнта старий браузер або вимкнений JavaScript, чи сайт зламається?»",
        "Ні, сайт спроєктовано за принципом прогресивного покращення (Progressive Enhancement). "
        "Вся базова розмітка, текстовий контент, характеристики станцій, каталог та початковий розрахунок калькулятора "
        "статично згенеровані безпосередньо в HTML. Якщо JS вимкнено, користувач все одно бачить повноцінний сайт, контакти та станції. "
        "JavaScript лише додає динамічну інтерактивність: перерахунок формул і фільтрацію списку."
    )

    add_qa_card(
        "Т2",
        "«Чому в коді одночасно є і tailwind.min.css, і cdn.tailwindcss.com?»",
        "tailwind.min.css — це заздалегідь скомпільований файл стилів (11 КБ), який завантажується миттєво і усуває затримку рендерингу (FCP 0.6 с). "
        "Офіційний скрипт CDN ми підключили згідно з обов'язковою вимогою Етапу 3 лабораторної роботи. "
        "Але щоб CDN-парсер не блокував перший показ сторінки на мобільних пристроях, він підключається неблокуюче, у фоні."
    )

    add_qa_card(
        "Т3",
        "«Чому ви не використали React, Vue або Angular замість ванільного JS?»",
        "Для лендінгу та інформаційного агрегатора тягнути важкі SPA-фреймворки (150–300 КБ бандлу) — це надлишковий оверхед (overengineering), "
        "який критично погіршує швидкість завантаження (Core Web Vitals). Чистий ванільний JavaScript дозволив реалізувати весь функціонал "
        "калькулятора, фільтрів та меню у легких файлах загальним обсягом менше 20 КБ і забезпечити 0 мс блокування процесора."
    )

    # =========================================================================
    # ЗБЕРЕЖЕННЯ ТА КОНВЕРТАЦІЯ
    # =========================================================================
    print(f"Збереження пам'ятки DOCX: {DOCX_OUT}...")
    doc.save(DOCX_OUT)
    print("DOCX збережено успішно!")

def convert_to_pdf():
    print("Конвертація пам'ятки в PDF через Word COM...")
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        
        doc = word.Documents.Open(os.path.abspath(DOCX_OUT))
        doc.SaveAs(os.path.abspath(PDF_OUT), FileFormat=17) # 17 = wdFormatPDF
        doc.Close()
        word.Quit()
        pythoncom.CoUninitialize()
        print(f"PDF пам'ятки успішно створено: {PDF_OUT}")
    except Exception as e:
        print(f"Помилка експорту PDF: {e}")

if __name__ == "__main__":
    build_defense_doc()
    convert_to_pdf()
