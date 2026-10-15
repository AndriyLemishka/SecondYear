import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def generate_lab2_report():
    base_docx = os.path.join(os.path.dirname(__file__), 'лаб 2 .docx')
    if not os.path.exists(base_docx):
        base_docx = 'Лаб 2/лаб 2 .docx'
    
    doc = docx.Document(base_docx)

    # 1. Update Title Page Submitter
    # P23: Виконав:
    p23 = doc.paragraphs[23]
    p23.text = "Виконав:"
    p23.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p23.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    # P24: студент гр. ПП-24
    p24 = doc.paragraphs[24]
    p24.text = "студент гр. ПП-24"
    p24.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p24.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    # P25: Лемішка А.
    p25 = doc.paragraphs[25]
    p25.text = "Лемішка А."
    p25.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p25.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    # P26: Прийняла:
    p26 = doc.paragraphs[26]
    p26.text = "Прийняла:"
    p26.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p26.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    # P27: Спас Н.Я.
    p27 = doc.paragraphs[27]
    p27.text = "Спас Н.Я."
    p27.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p27.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    # P32: ЛЬВІВ – 2026
    p32 = doc.paragraphs[32]
    p32.text = "ЛЬВІВ – 2026"
    p32.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p32.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    # 2. Task 2 Introduction (P45)
    p45 = doc.paragraphs[45]
    p45.text = "У базі даних \"Сайт продажу 3D-моделей (маркетплейс)\" зберігається інформація про тривимірні цифрові ассети, їхніх авторів-продавців, покупців, тематичні категорії каталогу, ліцензійні умови використання, прикріплені файли геометрії, кошики замовлень, позиції замовлень та відгуки користувачів. Основними сутностями є: seller, buyer, category, license, model_3d, model_file, cart, cart_item, review."
    p45.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p45.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    # 3. Entity Descriptions (P46 to P54)
    # Note: original docx had 7 entities (P46-P52) for Zoo. We have 9 entities!
    # P46 to P52 will be overwritten, and 2 new paragraphs will be inserted before P53.
    task2_entities = [
        # 1. seller
        "Сутність seller містить атрибути id_seller, nickname, rating. Атрибут id_seller є ідентифікаційним номером продавця (3D-художника) та виконує роль первинного ключа. Атрибут nickname визначає нікнейм або псевдонім автора на маркетплейсі, а rating — динамічний числовий рейтинг довіри покупців.",
        # 2. buyer
        "Сутність buyer складається з атрибутів id_buyer, email, registration_date. Атрибут id_buyer є первинним ключем та унікально ідентифікує покупця. Атрибут email зберігає контактну адресу електронної пошти користувача, а registration_date — дату його реєстрації на платформі.",
        # 3. category
        "Сутність category включає атрибути id_category, name. Первинним ключем є id_category, який ідентифікує тематичний розділ онлайн-каталогу. Атрибут name зберігає назву категорії (наприклад, «Персонажі», «Архітектура», «Транспорт», «Low-poly» тощо).",
        # 4. license
        "Сутність license містить атрибути id_license, id_change, license_name, usage_type, allow_modification, creation_date, creation_time. Атрибут id_license виступає первинним ключем правового регламенту. Атрибут id_change є ідентифікатором зміни (версії) ліцензії, license_name — офіційною назвою ліцензії, usage_type — типом дозволеного використання (комерційне чи Editorial), allow_modification — ознакою дозволу на модифікацію сітки, creation_date — днем створення ліцензії, а creation_time — точним часом її створення.",
        # 5. model_3d
        "Сутність model_3d містить атрибути id_model, id_seller, id_category, id_license, name, description, price, polygon_count, has_texture. Атрибут id_model виступає первинним ключем цифрового товару. Атрибут name визначає комерційну назву моделі, description — її розгорнутий технічний опис, price — встановлену вартість, polygon_count — кількість полігонів у тривимірній сітці, has_texture — ознаку наявності текстурних карт. Атрибути id_seller, id_category та id_license є зовнішніми ключами, які посилаються на первинні ключі сутностей seller, category та license відповідно.",
        # 6. model_file
        "Сутність model_file складається з атрибутів id_file, id_model, format, size_mb. Первинним ключем є атрибут id_file, що ідентифікує конкретний файл геометрії. Атрибут format визначає формат цифрового файлу (FBX, OBJ, STL, BLEND), size_mb — обсяг файлу в мегабайтах. Атрибут id_model є зовнішнім ключем, що посилається на первинний ключ сутності model_3d.",
        # 7. cart
        "Сутність cart включає атрибути id_cart, id_buyer, status, total_amount, created_date. Первинним ключем є id_cart, який ідентифікує кошик замовлення. Атрибут status відображає поточний стан замовлення («Активний» або «Оплачений»), total_amount — загальну суму до оплати, created_date — дату формування кошика. Атрибут id_buyer є зовнішнім ключем і посилається на первинний ключ покупця із сутності buyer.",
        # 8. cart_item
        "Сутність cart_item складається з атрибутів id_item, id_cart, id_model, fixed_price. Атрибут id_item є первинним ключем окремої позиції у замовленні. Атрибут fixed_price фіксує вартість 3D-моделі на момент додавання до замовлення. Атрибути id_cart та id_model виступають зовнішніми ключами, що посилаються на первинні ключі сутностей cart та model_3d відповідно.",
        # 9. review
        "Сутність review містить атрибути id_review, id_model, id_buyer, rating, comment_text. Первинним ключем є id_review, що ідентифікує відгук користувача. Атрибут rating містить бальну оцінку товару (від 1 до 5), comment_text — текстовий коментар покупця щодо якості сітки й топології моделі. Атрибути id_model та id_buyer є зовнішніми ключами, які посилаються на первинні ключі відповідних сутностей model_3d та buyer."
    ]

    # Update first 7 entities into paragraphs 46..52
    for i in range(7):
        p = doc.paragraphs[46 + i]
        p.text = task2_entities[i]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

    # Insert remaining 2 entities (cart_item, review) before paragraph 53
    p53_orig = doc.paragraphs[53]
    for text in task2_entities[7:]:
        p_new = p53_orig.insert_paragraph_before()
        p_new.text = text
        p_new.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p_new.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

    # Now find the paragraph: "Модель таблиць БД..."
    for p in doc.paragraphs:
        if "Модель таблиць БД" in p.text and "Session" not in p.text:
            p.text = "Модель таблиць БД «Сайт продажу 3D-моделей (маркетплейс)» зображено на рис. 2."
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
            break

    # Replace image 2 (Task 2 diagram)
    # Find the paragraph with drawing after the model announcement
    for i, p in enumerate(doc.paragraphs):
        if "Модель таблиць БД «Сайт продажу 3D-моделей" in p.text:
            # Next paragraph is the image paragraph
            p_img = doc.paragraphs[i + 1]
            p_img.text = "" # clear existing image
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_img = p_img.add_run()
            diag_path = os.path.join(os.path.dirname(__file__), 'DbDesigner_Marketplace.png')
            if not os.path.exists(diag_path):
                diag_path = 'Лаб 2/DbDesigner_Marketplace.png'
            r_img.add_picture(diag_path, width=Inches(6.4))
            
            # Next paragraph is the caption
            p_cap = doc.paragraphs[i + 2]
            p_cap.text = "Рис. 2. Модель таблиць БД «Сайт продажу 3D-моделей (маркетплейс)»"
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p_cap.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
            break

    # 4. Update Conclusions
    conclusions = [
        "Висновки: У ході виконання лабораторної роботи було закріплено теоретичні знання та практичні навички щодо логічного та фізичного моделювання реляційних баз даних. На основі детального опису предметних областей було побудовано фізичну модель навчальної БД «Session» у середовищі DbDesigner, а також спроєктовано цілісну фізичну модель бази даних для власної предметної області «Сайт продажу 3D-моделей (маркетплейс)», обраної в лабораторній роботі № 1.",
        "У процесі проєктування було реалізовано концептуальну схему у вигляді набору взаємопов'язаних таблиць, визначено коректні типи даних стовпців (integer, varchar, decimal, date, time), сформовано первинні ключі для однозначної ідентифікації сутностей та налаштовано зовнішні ключі для забезпечення реляційної цілісності зв'язків типу 1:N між сутностями seller, buyer, category, license, model_3d, model_file, cart, cart_item, review. Зокрема, у сутності license було успішно збережено розширені параметри версійності та часу створення (id_change, creation_date, creation_time).",
        "Практична робота в онлайн-середовищі DbDesigner наочно продемонструвала фундаментальні засади переходу від концептуального сприйняття предметної області до конкретної фізичної структури реляційних таблиць, готових до безпосереднього розгортання в сучасних реляційних системах керування базами даних (СУБД)."
    ]

    # Find where conclusions start
    conc_start_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.startswith("Висновки:"):
            conc_start_idx = idx
            break

    if conc_start_idx is not None:
        p_c0 = doc.paragraphs[conc_start_idx]
        p_c0.text = ""
        r_head = p_c0.add_run("Висновки: ")
        r_head.bold = True
        r_head.font.name = 'Times New Roman'
        r_head.font.size = Pt(12)
        r_body = p_c0.add_run(conclusions[0].replace("Висновки: ", ""))
        r_body.font.name = 'Times New Roman'
        r_body.font.size = Pt(12)
        p_c0.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        p_c1 = doc.paragraphs[conc_start_idx + 1]
        p_c1.text = conclusions[1]
        p_c1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p_c1.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

        p_c2 = doc.paragraphs[conc_start_idx + 2]
        p_c2.text = conclusions[2]
        p_c2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p_c2.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

    # 5. Save Document
    save_targets = [
        os.path.join(os.path.dirname(__file__), "Лемішка_А_ПП-24_Лаб2.docx"),
        os.path.join(os.path.dirname(__file__), "Лемішка А_ПП-24_Лаб2.docx")
    ]
    for out_name in save_targets:
        try:
            doc.save(out_name)
            print(f"Report successfully generated and saved to: {out_name}")
        except Exception as e:
            print(f"Notice: Could not save to {out_name}: {e}")

if __name__ == "__main__":
    generate_lab2_report()
