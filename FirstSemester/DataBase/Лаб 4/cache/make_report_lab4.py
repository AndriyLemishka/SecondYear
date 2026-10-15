import os
import json
import docx
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT

base_dir = r"d:\Studing\SecondYear\FirstSemester\DataBase\Лаб 4"
cache_dir = os.path.join(base_dir, "cache")
results_path = os.path.join(cache_dir, "query_results.json")
logo_path = os.path.join(cache_dir, "logo_lp.png")
target_docx = os.path.join(base_dir, "Лемішка А_ПП-24_ЛР4.docx")

with open(results_path, "r", encoding="utf-8") as f:
    query_data = json.load(f)

doc = docx.Document()

# Margins: 2 cm top/bottom, 2.5 cm left, 1.5 cm right
for s in doc.sections:
    s.top_margin = Cm(2)
    s.bottom_margin = Cm(2)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(1.5)

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(14)
font.color.rgb = RGBColor(0, 0, 0)

def add_p(text="", bold=False, italic=False, align='JUSTIFY', space_after=Pt(4), space_before=Pt(0), indent=Cm(1.25), size=Pt(14)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.line_spacing = 1.15
    if align == 'CENTER':
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.paragraph_format.first_line_indent = 0
    elif align == 'RIGHT':
        p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        p.paragraph_format.first_line_indent = 0
    elif align == 'LEFT':
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        p.paragraph_format.first_line_indent = indent
    else:
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
        p.paragraph_format.first_line_indent = indent
        
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = size
    return p

def add_code(text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(20, 20, 80)
    return p

def add_fig(img_path, caption, width=Inches(6.2)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_img.paragraph_format.first_line_indent = 0
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        run = p_img.add_run()
        run.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_cap.paragraph_format.first_line_indent = 0
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(8)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11.5)
        run_cap.italic = True
    else:
        add_p(f"[{caption}]", align='CENTER', indent=0)

print("Assembling Title Page...")

# ==========================================
# 1. ТИТУЛЬНИЙ ЛИСТ
# ==========================================
add_p("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", align='CENTER', size=Pt(13))
add_p("Національний університет “Львівська політехніка”", align='CENTER', size=Pt(13))
add_p("Інститут комп'ютерних наук та інформаційних технологій", align='CENTER', size=Pt(13))
add_p("Кафедра систем автоматизованого проектування", align='CENTER', space_after=Pt(16), size=Pt(13))

# Logo
if os.path.exists(logo_path):
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p_logo.paragraph_format.first_line_indent = 0
    p_logo.paragraph_format.space_before = Pt(4)
    p_logo.paragraph_format.space_after = Pt(16)
    p_logo.add_run().add_picture(logo_path, width=Inches(1.8))

add_p("Звіт", bold=True, align='CENTER', size=Pt(18))
add_p("до лабораторної роботи № 4", align='CENTER', size=Pt(15))
add_p("з курсу", align='CENTER', size=Pt(13))
add_p("«Організація баз даних та знань»", bold=True, align='CENTER', size=Pt(15))
add_p("на тему: “Написання запитів SELECT. Ч. 1”", bold=True, align='CENTER', space_after=Pt(40), size=Pt(14))

# Student / Teacher block
p_exec = doc.add_paragraph()
p_exec.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
p_exec.paragraph_format.left_indent = Cm(9.5)
p_exec.paragraph_format.first_line_indent = 0
p_exec.paragraph_format.line_spacing = 1.25
r_exec = p_exec.add_run("Виконав:\nстудент гр. ПП-24\nЛемішка А.А.\n\nПрийняла:\nСпас Н.Я.")
r_exec.font.name = 'Times New Roman'
r_exec.font.size = Pt(14)

add_p("ЛЬВІВ – 2026", align='CENTER', space_before=Pt(50), space_after=Pt(12), indent=0)
doc.add_page_break()

# ==========================================
# 2. МЕТА ТА ВСТУП
# ==========================================
p_meta = doc.add_paragraph()
p_meta.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
p_meta.paragraph_format.first_line_indent = Cm(1.25)
p_meta.paragraph_format.space_after = Pt(8)
r_m1 = p_meta.add_run("Мета роботи: ")
r_m1.bold = True
r_m1.font.name = 'Times New Roman'
r_m1.font.size = Pt(14)
r_m2 = p_meta.add_run("отримання практичних навичок створення SQL-запитів на вибірку, сортування, фільтрування та групування даних, застосування агрегатних функцій, а також створення та використання представлень (VIEW) у реляційній СУБД MySQL для бази даних індивідуального завдання «Сайт продажу 3D-моделей (маркетплейс)» (marketplace_3d).")
r_m2.font.name = 'Times New Roman'
r_m2.font.size = Pt(14)

add_p("Коротка характеристика предметної області та БД", bold=True, align='LEFT')
add_p("База даних marketplace_3d спроєктована для обслуговування онлайн-маркетплейсу цифрових 3D-моделей. Вона містить таблиці: seller (автори-продавці 3D-робіт), buyer (покупці контенту), category (тематичні категорії моделей), license (умови ліцензування), model_3d (каталог 3D-моделей із цінами та технічними параметрами), model_file (файли форматів FBX, OBJ, BLEND), cart (кошики/замовлення покупців), cart_item (окремі позиції замовлення з фіксованою ціною) та review (оцінки та текстові відгуки клієнтів). Усі запити розроблені та виконані в середовищі MySQL Workbench.", align='JUSTIFY')

# ==========================================
# 3. ЗАВДАННЯ 1: ПРОСТА ВИБІРКА (5 запитів -> 5 скріншотів)
# ==========================================
add_p("Завдання 1. Написання запитів простої вибірки (SELECT, FROM, LIMIT, DISTINCT, IS NULL)", bold=True, align='LEFT', space_before=Pt(8))
add_p("Оператор SELECT є основою мови маніпулювання даними (DML). У межах завдання створено 5 SQL-сценаріїв для реалізації простої вибірки інформації з таблиць БД marketplace_3d з використанням операторів вибору полів, усунення дублікатів (DISTINCT), обмеження результуючого набору рядків (LIMIT) та перевірки наявності невизначених значень (IS NULL). До кожного запиту наведено скріншот результату виконання у MySQL Workbench.", align='JUSTIFY')

t1_items = query_data[0]["items"]

add_p("1.1. Запит вибірки каталогу 3D-моделей (базовий SELECT з вибіркою полів та LIMIT)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: вивести перелік перших 6 моделей каталогу із зазначенням ідентифікатора, назви, вартості, полігональності та наявності текстур.", align='JUSTIFY')
add_code(t1_items[0]["sql"])
add_fig(os.path.join(cache_dir, "shot_1_1.png"), "Рис. 1.1. Скріншот виконання запиту 1.1 у MySQL Workbench")

add_p("1.2. Запит вибірки унікальних форматів 3D-файлів (DISTINCT)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: отримати перелік неповторюваних розширень та форматів файлів (FBX, BLEND, OBJ), доступних для завантаження на майданчику.", align='JUSTIFY')
add_code(t1_items[1]["sql"])
add_fig(os.path.join(cache_dir, "shot_1_2.png"), "Рис. 1.2. Скріншот виконання запиту 1.2 у MySQL Workbench")

add_p("1.3. Отримання списку перших трьох зареєстрованих покупців (LIMIT)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: отримати інформацію про трьох найбільш ранніх покупців за датою реєстрації для маркетингової вибірки.", align='JUSTIFY')
add_code(t1_items[2]["sql"])
add_fig(os.path.join(cache_dir, "shot_1_3.png"), "Рис. 1.3. Скріншот виконання запиту 1.3 у MySQL Workbench")

add_p("1.4. Пошук моделей, у яких відсутній текстовий опис (IS NULL)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: виявити товари, для яких автор не вказав опис (значення поля description дорівнює NULL), для модерації каталогу.", align='JUSTIFY')
add_code(t1_items[3]["sql"])
add_fig(os.path.join(cache_dir, "shot_1_4.png"), "Рис. 1.4. Скріншот виконання запиту 1.4 у MySQL Workbench")

add_p("1.5. Вибірка унікальних статусів кошиків та сум зі зміщенням (LIMIT offset, count)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: демонстрація пагінації (посторінкового виведення) неповторюваних записів замовлень шляхом пропуску перших двох рядків та вибірки наступних чотирьох.", align='JUSTIFY')
add_code(t1_items[4]["sql"])
add_fig(os.path.join(cache_dir, "shot_1_5.png"), "Рис. 1.5. Скріншот виконання запиту 1.5 у MySQL Workbench")

# ==========================================
# 4. ЗАВДАННЯ 2: СОРТУВАННЯ ДАНИХ (3 запити -> 3 скріншоти)
# ==========================================
add_p("Завдання 2. Сортування даних у запитах (ORDER BY, DESC, розрахункові вирази)", bold=True, align='LEFT', space_before=Pt(8))
add_p("Оператор ORDER BY впорядковує результуючий набір записів за одним або кількома стовпцями. За замовчуванням застосовується порядок зростання (ASC), тоді як ключове слово DESC визначає сортування за спаданням. У цьому розділі реалізовано 3 запити сортування даних каталогу.", align='JUSTIFY')

t2_items = query_data[1]["items"]

add_p("2.1. Сортування авторів за рейтингом від найвищого до найнижчого (ORDER BY DESC)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: формування рейтингової таблиці найкращих продавців 3D-моделей для відображення у блоці «Топ автори».", align='JUSTIFY')
add_code(t2_items[0]["sql"])
add_fig(os.path.join(cache_dir, "shot_2_1.png"), "Рис. 2.1. Скріншот виконання запиту 2.1 у MySQL Workbench")

add_p("2.2. Сортування моделей за зростанням ціни з обчисленням вартості з ПДВ (20%) через псевдонім (AS)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: виведення товарів від найдешевших до найдорожчих з розрахунковим полем остаточної вартості з урахуванням податку.", align='JUSTIFY')
add_code(t2_items[1]["sql"])
add_fig(os.path.join(cache_dir, "shot_2_2.png"), "Рис. 2.2. Скріншот виконання запиту 2.2 у MySQL Workbench")

add_p("2.3. Сортування за двома стовпцями: наявність текстури (DESC) та кількість полігонів (ASC)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: пріоритетний показ моделей із наявними текстурами ('Yes' розміщується перед 'No'), а серед них — впорядкування за мірою зростання складності 3D-сітки (від оптимізованих low-poly до важких high-poly).", align='JUSTIFY')
add_code(t2_items[2]["sql"])
add_fig(os.path.join(cache_dir, "shot_2_3.png"), "Рис. 2.3. Скріншот виконання запиту 2.3 у MySQL Workbench")

# ==========================================
# 5. ЗАВДАННЯ 3: ФІЛЬТРАЦІЯ ДАНИХ (5 запитів -> 5 скріншотів)
# ==========================================
add_p("Завдання 3. Фільтрація даних за умовами (IN, BETWEEN, LIKE, REGEXP)", bold=True, align='LEFT', space_before=Pt(8))
add_p("Конструкція WHERE дає змогу вибирати лише ті записи, що задовольняють задані критерії. У межах завдання розроблено 5 запитів із застосуванням оператора перевірки належності до списку (IN), потрапляння у числовий діапазон (BETWEEN), маскового пошуку рядків (LIKE) та фільтрації за регулярними виразами (REGEXP).", align='JUSTIFY')

t3_items = query_data[2]["items"]

add_p("3.1. Фільтрація моделей за списком категорій (оператор IN)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: вибірка моделей, що належать виключно до трьох категорій: 'Characters' (id 1), 'Vehicles' (id 2) та 'Architecture' (id 8).", align='JUSTIFY')
add_code(t3_items[0]["sql"])
add_fig(os.path.join(cache_dir, "shot_3_1.png"), "Рис. 3.1. Скріншот виконання запиту 3.1 у MySQL Workbench")

add_p("3.2. Фільтрація файлів моделей за діапазоном розміру (оператор BETWEEN)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: пошук файлів моделей розміром від 10.0 до 50.0 мегабайт для оцінки середнього навантаження на сховище.", align='JUSTIFY')
add_code(t3_items[1]["sql"])
add_fig(os.path.join(cache_dir, "shot_3_2.png"), "Рис. 3.2. Скріншот виконання запиту 3.2 у MySQL Workbench")

add_p("3.3. Пошук моделей фантастичної тематики за шаблоном назви (оператор LIKE)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: пошук моделей, назва яких містить ключове слово 'Sci-Fi' у будь-якій позиції рядка за допомогою спецсимволів '%'.", align='JUSTIFY')
add_code(t3_items[2]["sql"])
add_fig(os.path.join(cache_dir, "shot_3_3.png"), "Рис. 3.3. Скріншот виконання запиту 3.3 у MySQL Workbench")

add_p("3.4. Фільтрація користувачів за маскою електронної пошти (оператор LIKE)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: знаходження клієнтів зі спеціалізованими акаунтами, що містять маркер '_art@' або стандартний префікс 'user_%'.", align='JUSTIFY')
add_code(t3_items[3]["sql"])
add_fig(os.path.join(cache_dir, "shot_3_4.png"), "Рис. 3.4. Скріншот виконання запиту 3.4 у MySQL Workbench")

add_p("3.5. Складна фільтрація назв моделей за регулярним виразом (оператор REGEXP)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: вибірка 3D-моделей, назва яких починається зі слів Cyberpunk, Sports або Space з використанням метасимволів ^ та альтернативи |.", align='JUSTIFY')
add_code(t3_items[4]["sql"])
add_fig(os.path.join(cache_dir, "shot_3_5.png"), "Рис. 3.5. Скріншот виконання запиту 3.5 у MySQL Workbench")

# ==========================================
# 6. ЗАВДАННЯ 4: АГРЕГАТНІ ФУНКЦІЇ (5 запитів -> 5 скріншотів)
# ==========================================
add_p("Завдання 4. Використання агрегатних функцій (COUNT, AVG, MIN, MAX, SUM)", bold=True, align='LEFT', space_before=Pt(8))
add_p("Агрегатні функції здійснюють обчислення над множиною значень стовпця та повертають єдине скалярне значення. У MySQL основними агрегатними функціями є COUNT, AVG, MIN, MAX та SUM. У цьому розділі створено 5 аналітичних запитів для розрахунку загальних показників системи.", align='JUSTIFY')

t4_items = query_data[3]["items"]

add_p("4.1. Підрахунок загальної кількості моделей та кількості категорій (COUNT, COUNT DISTINCT)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: визначення сумарного обсягу каталогу та кількості категорій, в яких вже розміщено товари.", align='JUSTIFY')
add_code(t4_items[0]["sql"])
add_fig(os.path.join(cache_dir, "shot_4_1.png"), "Рис. 4.1. Скріншот виконання запиту 4.1 у MySQL Workbench")

add_p("4.2. Розрахунок середньої вартості 3D-моделі та середньої кількості полігонів (AVG)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: оцінка середньоринкової вартості позиції в каталозі та середньої деталізації геометрії моделей.", align='JUSTIFY')
add_code(t4_items[1]["sql"])
add_fig(os.path.join(cache_dir, "shot_4_2.png"), "Рис. 4.2. Скріншот виконання запиту 4.2 у MySQL Workbench")

add_p("4.3. Визначення мінімальної та максимальної вартості товару в каталозі (MIN, MAX)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: визначення цінового діапазону (найдешевшої промо-моделі та найдорожчого професійного пакету).", align='JUSTIFY')
add_code(t4_items[2]["sql"])
add_fig(os.path.join(cache_dir, "shot_4_3.png"), "Рис. 4.3. Скріншот виконання запиту 4.3 у MySQL Workbench")

add_p("4.4. Обчислення загальної виручки за оплаченими замовленнями (SUM)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: розрахунок загального фінансового обороту маркетплейсу за кошиками у статусі 'Completed'.", align='JUSTIFY')
add_code(t4_items[3]["sql"])
add_fig(os.path.join(cache_dir, "shot_4_4.png"), "Рис. 4.4. Скріншот виконання запиту 4.4 у MySQL Workbench")

add_p("4.5. Комплексна статистика за параметрами файлів сховища (COUNT, MIN, MAX, AVG, SUM)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: отримання зведеного технічного звіту про кількість файлів, найменший, найбільший, середній обсяг та сумарний дисковий простір сховища.", align='JUSTIFY')
add_code(t4_items[4]["sql"])
add_fig(os.path.join(cache_dir, "shot_4_5.png"), "Рис. 4.5. Скріншот виконання запиту 4.5 у MySQL Workbench")

# ==========================================
# 7. ЗАВДАННЯ 5: ГРУПУВАННЯ ДАНИХ (3 запити -> 3 скріншоти)
# ==========================================
add_p("Завдання 5. Групування вибраних даних та фільтрація груп (GROUP BY, HAVING)", bold=True, align='LEFT', space_before=Pt(8))
add_p("Оператор GROUP BY об'єднує рядки з однаковими значеннями у групи для застосування агрегатних функцій окремо до кожної підгрупи. Для фільтрації результатів групування використовується конструкція HAVING, яка на відміну від WHERE працює зі значеннями обчислених агрегатів.", align='JUSTIFY')

t5_items = query_data[4]["items"]

add_p("5.1. Розподіл кількості моделей та їх середньої вартості за категоріями (GROUP BY)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: аналіз товарного наповнення кожної окремої категорії та визначення середнього цінового чеку за напрямками.", align='JUSTIFY')
add_code(t5_items[0]["sql"])
add_fig(os.path.join(cache_dir, "shot_5_1.png"), "Рис. 5.1. Скріншот виконання запиту 5.1 у MySQL Workbench")

add_p("5.2. Аналіз статусів замовлень із сумарним фінансовим обсягом понад 50 (GROUP BY, HAVING)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: виділити тільки ті статуси кошиків, сумарний оборот за якими перевищує встановлену межу в 50 у.о.", align='JUSTIFY')
add_code(t5_items[1]["sql"])
add_fig(os.path.join(cache_dir, "shot_5_2.png"), "Рис. 5.2. Скріншот виконання запиту 5.2 у MySQL Workbench")

add_p("5.3. Визначення авторів з двома і більше моделями та вартості їх портфоліо (GROUP BY, HAVING)", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: виявлення постійних та продуктивних продавців платформи (кількість робіт >= 2) та оцінка загальної вартості їхнього контенту.", align='JUSTIFY')
add_code(t5_items[2]["sql"])
add_fig(os.path.join(cache_dir, "shot_5_3.png"), "Рис. 5.3. Скріншот виконання запиту 5.3 у MySQL Workbench")

# ==========================================
# 8. ЗАВДАННЯ 6: ПРЕДСТАВЛЕННЯ (VIEW) (2 представлення -> 2 скріншоти)
# ==========================================
add_p("Завдання 6. Створення та використання представлень (CREATE VIEW, SELECT FROM VIEW)", bold=True, align='LEFT', space_before=Pt(8))
add_p("Представлення (VIEW) — це віртуальна таблиця, вміст якої визначається збереженим на сервері SQL-запитом. Представлення не дублюють фізичні дані, а динамічно формують результат під час звернення. Вони спрощують розробку складних запитів, інкапсулюють об'єднання кількох таблиць та підвищують безпеку доступу. У межах роботи створено два представлення.", align='JUSTIFY')

t6_items = query_data[5]["items"]

add_p("6.1. Створення та тестування представлення view_popular_models", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: формування денормалізованої вітрини товарів із підстановкою реальної назви категорії та авторського нікнейму продавця замість зовнішніх числових ключів (id_category, id_seller).", align='JUSTIFY')
add_p("Скрипт створення представлення view_popular_models:", bold=True, indent=Cm(0.6))
add_code(t6_items[0]["create_sql"])
add_p("Запит вибірки з представлення (моделі від 20 у.о. з сортуванням за спаданням ціни):", bold=True, indent=Cm(0.6))
add_code(t6_items[0]["select_sql"])
add_fig(os.path.join(cache_dir, "shot_6_1.png"), "Рис. 6.1. Скріншот створення та вибірки з представлення view_popular_models у MySQL Workbench")

add_p("6.2. Створення та тестування представлення view_buyer_order_summary", bold=True, align='LEFT', indent=Cm(0.6))
add_p("Призначення: формування зведеної аналітики активності покупців із підрахунком кількості завершених замовлень та загальної витраченої суми грошей (із використанням LEFT JOIN та COALESCE).", align='JUSTIFY')
add_p("Скрипт створення представлення view_buyer_order_summary:", bold=True, indent=Cm(0.6))
add_code(t6_items[1]["create_sql"])
add_p("Запит вибірки з представлення (активні покупці з покупками, відсортовані за спаданням суми витрат):", bold=True, indent=Cm(0.6))
add_code(t6_items[1]["select_sql"])
add_fig(os.path.join(cache_dir, "shot_6_2.png"), "Рис. 6.2. Скріншот створення та вибірки з представлення view_buyer_order_summary у MySQL Workbench")

# ==========================================
# 9. ВІДПОВІДІ НА КОНТРОЛЬНІ ЗАПИТАННЯ
# ==========================================
add_p("Відповіді на типові контрольні запитання", bold=True, align='LEFT', space_before=Pt(12))

questions_answers = [
    (
        "1. За допомогою якого запиту можна вивести інформацію з певної таблиці БД?",
        "Для виведення інформації з таблиці реляційної бази даних використовується базовий оператор мови DML SELECT. Синтаксис: SELECT <список_стовпців | *> FROM <ім'я_таблиці>;. Використання зірочки (*) забезпечує вибірку всіх стовпців таблиці, а перерахування конкретних імен через кому оптимізує передачу даних та дозволяє вибирати тільки потрібні поля. За потреби запит доповнюється конструкціями WHERE, ORDER BY, LIMIT тощо."
    ),
    (
        "2. Який запит дає змогу вибрати унікальні значення виробників із БД?",
        "Для вибірки неповторюваних значень використовується ключове слово DISTINCT після оператора SELECT. Наприклад, для таблиці товарів: SELECT DISTINCT manufacturer FROM products;. У разі наявності декількох однакових записів виробника у вихідній таблиці, у фінальний результат кожне ім'я потрапить рівно один раз. Для нашої БД аналогічний запит: SELECT DISTINCT format FROM model_file;."
    ),
    (
        "3. Назвіть оператори фільтрації та охарактеризуйте їх застосування?",
        "Фільтрація рядків здійснюється за допомогою розділу WHERE. Основні оператори:\n"
        "• Оператори порівняння (=, !=, <>, <, >, <=, >=) — перевірка точної відповідності або числових меж.\n"
        "• IN (val1, val2, ...) / NOT IN — перевірка входження значення у заданий список або результат підзапиту.\n"
        "• BETWEEN min AND max — перевірка належності значення до замкненого числового чи датового діапазону.\n"
        "• LIKE / NOT LIKE — пошук рядків за спрощеним шаблоном із підстановками '%' та '_'.\n"
        "• REGEXP / RLIKE — розширена фільтрація за регулярними виразами.\n"
        "• IS NULL / IS NOT NULL — перевірка на наявність або відсутність порожнього/невизначеного значення NULL.\n"
        "• Логічні оператори AND, OR, NOT — для побудови складених умов фільтрації."
    ),
    (
        "4. Запишіть синтаксис оператора сортування стовпців.",
        "Синтаксис оператора сортування в SQL має такий вигляд:\n"
        "SELECT <стовпці> FROM <таблиця> [WHERE умова] ORDER BY стовпець_1 [ASC | DESC] [, стовпець_2 [ASC | DESC], ...];\n"
        "Ключове слово ASC визначає порядок сортування за зростанням (встановлено за замовчуванням), а DESC — за спаданням."
    ),
    (
        "5. Як задати псевдонім стовпця?",
        "Псевдонім (аліас) стовпця задається за допомогою ключового слова AS після імені поля або обчислювального виразу: SELECT price * 1.20 AS price_with_vat FROM model_3d;. Використання AS є необов'язковим (можна просто вказати нове ім'я через пробіл), проте вважається хорошим тоном для читабельності коду. Якщо псевдонім містить пробіли, його беруть у зворотні лапки (backticks) або одинарні лапки."
    ),
    (
        "6. Символи та їх зміст в операторах пошуку за шаблоном.",
        "В операторі LIKE використовуються два спеціальні символи:\n"
        "• '%' (відсоток) — відповідає будь-якій кількості довільних символів (включаючи нульову кількість);\n"
        "• '_' (підкреслення) — відповідає рівно одному будь-якому символу.\n"
        "В операторі REGEXP використовуються метасимволи регулярних виразів: '^' (початок рядка), '$' (кінець рядка), '.' (будь-який одиночний символ), '[...]' (будь-який символ із набору), '|' (логічне АБО), '*' та '+' (повторення)."
    ),
    (
        "7. Який запит дає змогу сортувати стовпці за спаданням?",
        "Сортування за спаданням здійснюється додаванням ключового слова DESC після назви сортованого стовпця в реченні ORDER BY. Наприклад: SELECT id_seller, nickname, rating FROM seller ORDER BY rating DESC; — сортує продавців від найбільшого рейтингу до найменшого."
    ),
    (
        "8. Запишіть запит сортування за декількома стовпцями у таблиці БД.",
        "Приклад запиту багаторівневого сортування: SELECT id_model, name, has_texture, polygon_count, price FROM model_3d ORDER BY has_texture DESC, polygon_count ASC;. Рядки спершу впорядковуються за наявністю текстури (значення 'Yes' будуть вище за 'No'), а для записів з однаковим значенням текстури застосовується друге правило — зростання кількості полігонів."
    ),
    (
        "9. Які агрегатні функції використовують у запитах?",
        "У СУБД MySQL застосовуються такі стандартні агрегатні функції:\n"
        "• COUNT(вираз / *) — повертає загальну кількість рядків або непорожніх (non-NULL) значень;\n"
        "• SUM(стовпець) — обчислює арифметичну суму числових значень;\n"
        "• AVG(стовпець) — обчислює середнє арифметичне значення у числовому стовпці;\n"
        "• MIN(стовпець) — знаходить найменше значення (для чисел, дат чи рядків);\n"
        "• MAX(стовпець) — знаходить найбільше значення."
    ),
    (
        "10. За допомогою яких операторів можна групувати дані?",
        "Дані групуються за допомогою оператора GROUP BY, після якого перелічуються один або кілька стовпців групування: SELECT id_category, COUNT(*) FROM model_3d GROUP BY id_category;. Усі рядки, що мають однакові значення в зазначених стовпцях, об'єднуються в одну підгрупу, над якою обчислюються агрегатні функції."
    ),
    (
        "11. Як здійснити фільтрацію груп?",
        "Фільтрація сформованих груп здійснюється за допомогою оператора HAVING, який розташовується після GROUP BY: SELECT id_category, COUNT(*) FROM model_3d GROUP BY id_category HAVING COUNT(*) >= 2;. На відміну від WHERE (який фільтрує окремі рядки до їх агрегації), HAVING відсікає цілі групи на основі результатів агрегатних виразів."
    ),
    (
        "12. Для чого використовують представлення (VIEW)?",
        "Представлення (View) використовуються для:\n"
        "1) Спрощення складних запитів: інкапсуляція багаторазових з'єднань JOIN та підсумкових обчислень в єдиний іменований об'єкт;\n"
        "2) Безпеки даних: обмеження доступу користувачів до певних колонок чи рядків без надання прав на повні базові таблиці;\n"
        "3) Логічної незалежності: можливість змінювати внутрішню структуру фізичних таблиць без порушення роботи зовнішніх клієнтських програм."
    ),
    (
        "13. Що таке віртуальна таблиця?",
        "Віртуальна таблиця (представлення, View) — це іменований об'єкт бази даних, що являє собою збережений на сервері SQL-запит вибірки. Віртуальна таблиця не зберігає власних фізичних копій даних на диску (за винятком матеріалізованих представлень). Під час кожного звернення до віртуальної таблиці через оператор SELECT СУБД динамічно виконує базовий запит та формує актуальний набір даних із реальних таблиць."
    )
]

for q_text, a_text in questions_answers:
    p_q = doc.add_paragraph()
    p_q.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    p_q.paragraph_format.first_line_indent = Cm(1.25)
    p_q.paragraph_format.space_before = Pt(4)
    p_q.paragraph_format.space_after = Pt(2)
    r_q = p_q.add_run(q_text)
    r_q.bold = True
    r_q.font.name = 'Times New Roman'
    r_q.font.size = Pt(13)

    p_a = doc.add_paragraph()
    p_a.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    p_a.paragraph_format.first_line_indent = Cm(1.25)
    p_a.paragraph_format.space_before = Pt(0)
    p_a.paragraph_format.space_after = Pt(4)
    r_a = p_a.add_run(a_text)
    r_a.font.name = 'Times New Roman'
    r_a.font.size = Pt(13)

# ==========================================
# 10. ВИСНОВОК (РІВНО 1 АБЗАЦ!)
# ==========================================
add_p("Висновки", bold=True, align='CENTER', space_before=Pt(12), space_after=Pt(6), indent=0)

conclusion_single_paragraph = (
    "Під час виконання лабораторної роботи № 4 було успішно досягнуто поставленої мети та здобуто комплексні "
    "практичні навички роботи з мовою SQL у СУБД MySQL Workbench для бази даних маркетплейсу 3D-моделей marketplace_3d: "
    "на практиці опановано побудову запитів простої вибірки за допомогою оператора SELECT із обмеженням кількості "
    "записів LIMIT, виключенням дублікатів DISTINCT та пошуком невизначених значень IS NULL; реалізовано механізми "
    "одно- та багаторівневого сортування даних за допомогою ORDER BY, модифікаторів ASC/DESC та розрахункових виразів; "
    "засвоєно фільтрацію рядків операторами IN, BETWEEN, LIKE (із символами % та _) і регулярними виразами REGEXP; "
    "застосовано агрегатні функції COUNT, AVG, MIN, MAX, SUM для розрахунку ключових бізнес-показників; опановано "
    "групування даних оператором GROUP BY та фільтрацію агрегованих груп через HAVING; а також створено та протестовано "
    "віртуальні таблиці (VIEW) для формування зручної вітрини товарів і зведеної аналітики замовлень клієнтів."
)

add_p(conclusion_single_paragraph, align='JUSTIFY', space_before=Pt(2), space_after=Pt(6), indent=Cm(1.25))

# Saving DOCX
doc.save(target_docx)
print(f"Report document successfully updated: {target_docx}")
